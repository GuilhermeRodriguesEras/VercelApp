from flask import Flask, request, jsonify, redirect
from flask_cors import CORS

import requests
import os
import json
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
import secrets
import time
from hashlib import sha256
from urllib.parse import urlencode

import base64
from io import BytesIO
from html import unescape, escape as html_escape
from html.parser import HTMLParser

from flask import send_file

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

class TinyHTMLParagraphParser(HTMLParser):

    TAGS_FORMATACAO = {"b", "strong", "i", "em", "u"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.partes = []

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag == "br":
            self.partes.append("<br/>")
        elif tag in self.TAGS_FORMATACAO:
            tag_saida = "b" if tag == "strong" else "i" if tag == "em" else tag
            self.partes.append(f"<{tag_saida}>")
        elif tag in {"p", "div"}:
            pass
        elif tag == "li":
            self.partes.append("- ")

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in self.TAGS_FORMATACAO:
            tag_saida = "b" if tag == "strong" else "i" if tag == "em" else tag
            self.partes.append(f"</{tag_saida}>")
        elif tag in {"p", "div", "li"}:
            self.partes.append("<br/>")

    def handle_data(self, data):
        if data:
            self.partes.append(html_escape(data).replace("\n", "<br/>"))

    def resultado(self):
        import re
        resultado = "".join(self.partes)
        resultado = re.sub(r"(?:<br/>\s*){3,}", "<br/><br/>", resultado)
        return resultado.strip("<br/> \n\t")


def html_tiny_para_reportlab(valor):
    if not valor:
        return ""

    parser = TinyHTMLParagraphParser()
    try:
        parser.feed(unescape(str(valor)))
        parser.close()
        resultado = parser.resultado()
        return resultado or html_escape(unescape(str(valor)))
    except Exception:
        return html_escape(unescape(str(valor))).replace("\n", "<br/>")

app = Flask(__name__)

BRFER_LOGO_BASE64 = "iVBORw0KGgoAAAANSUhEUgAAAJUAAABwCAIAAACYSaUpAAAACXBIWXMAAA7EAAAOxAGVKw4bAAABZGlDQ1BJQ0NCYXNlZChSR0IsR29vZ2xlL1NraWEvN0M1RkEyMTUxMzk3NDc0QTA0ODZCQkNDODM3MzNENTkpAAB4nH2QvUrDYBSGH2tBFMVBhw4OGRxc1P5of8ClrVhcW4VWpzRNi9ifkKboBejm4OomLt6A6GUoCA7i4CWIoLNvGiQFqefw5nt485Iv50Akhioah07Xc8ulglGtHRhT70yoh2VafYfxpdT3S5B9Xv0nN66mG3bf0vkhea4u1ycb4sVWwKc+1wO+8PnEczzxtc/uXrkovhOvtEa4PsKW4/r5N/FWpz2wwv9m1u7uV3RWpSVK9NQt2tisU+GYI0xRhiKb7JAnSUKUIEVO7sZQeeJ6ZklTUBfVWb3PSCm2lc75+wyu7N1A9gsmL0OvfgUP5xB7Db1lzTZ/BvePoRfu2DFdc2hFpUizCZ+3MFeDhSeYOfxd7JhZjT+zGuzSxWJNlNQ0CdI/hc1LvY60eocAACraSURBVHic7X0HXFTH9r8v7728kvfy0qOCSi/L0pcuXVEQQRAQUZAmilJEpQlRQQUUsTew946IMfaoSSxRY9SY2Cv2xK4oZe/9nd3ZnZ29bXcBTd7//85nP8vduWfOnJnvnDNzZuZe2lFtR1KplJFC07TAT3YiJwP7Llywy2Jzsr9/LxLWliSkJ/CTlRXgb9dKzVpG2tenDSVI5dTKct8CMZTUgB8tJ8T3RolRBPnzDZVONoGwPozr1ujDFqWrNLZiAhLa0Vq7L12JTxVh/jYsmpHCINzNST3bRCVOUQyBnAza6MDQViv/2eJqtDgvJ7WV9xNo39+RUO1Ixdgp7CztBIz3/x9qw1qjlm0raRqJiR+DSLRxCuMu+VPLUil1T6JNuQISKJZJsW/xSRC+ptWdlUChnArwlcjJz6i7AANJb9B/sjMKyGGoKyxNe30YgxxDFK0pDmlz0r6OjBS+jBrsj01YaBuSlgLbvNw/GrWggsz4QUAEpW4BrSdd5fDx66RSC/RnNw6n3QgrqT2DTtRC/8lWiFE39jUjha/ywgpwasIuRaMoWr0dBZRnF4Gu2VG2sCgBzQWqQxHLDpz1+n3WXyit6/PWSv+DkK6KaTv+YeltS29C5u9LrWmoFmRUsz+2CL7eobFUYVUodcfCSGcXqpGHkci41qgDX9Vo+QSVzcaZXeCaTBHgYUMg0A6IdPOf7Nq2LDty6MJzdzaPcOl8+DF4hBc1tFGpBURm5BPC1zXZQJKkQ/xA8fflllEbCmylKJ2yt3k7tIZU8YOwWioGdS5K/pHSOLNU+VHewykUTlHJFC5R+2pQLC/UVkRaAF9xArdaU6I2EtTw4yN1Bt4xUiCFU0u5J1Egz+k6eBTg5SHrrCUbTQDA4OHLy1lZja2vUX+GGuwsnHJ0jh943bdU9lHw0LKPnJoRQhxyaLj3u+2mat+giH6vjV+N5SrwEzAUBklpHmbCPaq7SUFh8g/RS7VX5A9HOjVjWxHTf2rMICUh4oOSNZvCP2VWJ/tQxEipHC9p5uhIkkbFGGw6Zad4nB77p8ZEdrnCCqg1Dn9BfAy84x8pgiTkGFWWo7JvBTAChPjlWWTMxE8s4Y0Tzeqsb0j42yHN8QPF6I/YVmiGFRJmxCT1+SdFs+ao/yMFUYLmyKaW4Mduc3Z5HBqoZeNGmk9r7etD00w/JsyGOTWyaamMlmJ1JT5RqvNLtGpOz9EFcHOguSUpFM82yVyMJXOaUutZuGllH2KCKn0z7oezQclac+jGcoaMn5zpjGvGiUVGS7LTGQzCKqHrluw/yJYEleMWAg+3uyyxWcqppbwuTeRwR1HyEII/4mkNvc0ZP98uPztR1zpqjh+0N2HSmBihgrLnKGIBCndz5jSVGPakjESp4Aj6P+ImtfOfFOFe2KS4iyxMfaKPkKNk5tXYTDU00k0KjyoHj1J3BXLblSOlhjcTP3YvFtaNwcno8tpnpwgr4cvLeYvTPkhRfEULaMiWwLhg4sduL2a6PGygVGEcLfeCTfDdTL+up54WVxRu2L62iW6kiGEPLbUgx4ukKAZIFEIoStEcgbSe2I1CtenJUo08GstiaEimsxN19p9KSBSz0Nev6787cuDqrfPnbp6Zu2p6xJBezoFWgVF+r5pfKsIEhQcF6JoR2A3Uq2NnjmzdVS2lG2VilH74f9QCagl+JAEG62uXe/aR9Ej09hjg4BZt4xommjgrv4l+JYvQUe+mG5voBqnMIsHzNjxquB+REpJdmtVAv5DxKPCTsoUrCiR6nw5qtihX64lSd13sa5200sjM9J+M+IHLbJtl00g0sIE7pBsu3vsxLL2byyArx3hLp4HmHlHWs5eWNtP1yB8C8607N3Yf+GrfgV2/PbxXd/9a1fo5bmH2AXFevzXcRSYox1Xujd9MAIGrgDfTaXX/SXF5LUbL0Dy+EUsWzk7xOECKhStDAp9MdN1O2WRo4KGRENSgZGBAyUOF5y/qH7++e/T0gf1HdsNURYYO3XjjyaXhk+LdYqxcBlo4xZi59rPqPdj/zuurdxtu1B7eXFA14osFme59LV3CTV2jLV2ixG5RIucoU0m0Rf6c7LpH9yjZRKaRpp/RdD0aDqVSxcxGiroIRTfRVBONMUbKKPDma1aFznLA3loswYmQMKicQiitnyJD/lN94k6Rf/Hpj+bnLx57+bqKXU1s3E16R/lfuHb6/tNbXx7c0mdoT48YO/cBYpd+ZoCQR7S4a5Q9JIYP7+UT4+Ie5+Ad5+Aaae4UaebaX+wUYeEi+5h5RFqEJfrcvHNOjgcghIZHWrm61kR4VNlPpImAN0H14ay2AD87Cye/cCPytThnCn42U6Bctg58irVTrEMrJ354QoFhU8YAzTduXhw3YbSJo76lSyexR2fPYFuvEHtJsMghxNQjxsYlSmQfbAgfh94mkAIfx3AZnG6D7DwG2TpEmblEW8JPwM81EmzR0r+fpGpt+aPn18ADy8IOShZzKN0bwq9J7n7RWkEzxo+sBL4Gtrt37/6gTidOnPjpp5/u37/PCeT58+cR20k5oevjx49fuHDh9evXZIvD99OnTzEziEX8OGNdXR3C4+zZs7hoLBP9PHXqVHOzrIs+efKEfRe+T58+fe/ePQaWFMt8OfFTd/Ryg6DwPjtafKaoI8cO2jqamEg6mrt0tHRvb+PT0T6gi02PLuLATuASXSOsHAG2YFP4tg8xkoQZSyKMXGPMXeOsXOPFboPErgPFbv3Bc5pL+pp3jbJdsGX20C+GjSgeOWXe2KdPTtNNF2lpHUU/pKQNCi2kUoUbV4SXyj6o3PrAoybSs6Kioh0X/elPfzIxMamqqiJPksG3h4cHusvm/89//pORkVFfX4/b5MCBA/guu4jc3FzEZmVlxRCFLz788MPGRllMtXfvXk49EZuDgwMwCNg6g7jnn2pQK0ZE2TRy3KQCU8cOgJ/Io4Otrx5AaOOvbxdg4BRiBiYlCTO1723iGGrm2MfEJcLENdrMNcYC8HOKtXIZZO0+SOw+UOTU30oSLQpL8x5RMsinv4d7X6e0HN87v4w9vyv06qHU5/dX0c03aelTvKeBTK+Z4Tkpmh3vA36cYODr8ePHk/VC+GEedt6BAwdi/v379/OBB4mc+DGYEX4gCuPHp+17770HxsppbZz2x4UfC0tKNgQ1Ll6xEOFn7dXZ2lff2k/Ppltn+x6Gkt6mzmEmTuEmDuHGDmGmMNSBqbn2FwFU7nE2kgE2zrH2LrEilzhLyQCxU3/LqGHWG7cNW7EuflNt2rFDqWe+tLm1rf2lGtOjm7se2VX48PbXNPWUkkf9smJlO4XNspmMMuhUqqe2DETaH4aEbKN3330XHB0bPz7I//znP1+6dAlV/+uvv8YM77zzDiMLn/2R0j766KOmpiYSPwHCXYfhSLnxY9okjVdF5OOfvI9LZXPEhvFlX5g5dzBz+dzatzOYnW0QDHjGjsEmklAZeJIIU0mEsWukqXs/c9docxjw3GKs3WMsnQfYucVLfFPseqSKe6fZxY5yzCu0OFDb/Vit67GtHkc22v2y1fTylvY/bzI4usHuq7Xhl39ZdOvej5fv3XpJ040IJ2mj7CNfuFPMS5W+FGs9bdo0XP8+ffosXrx44cKFeXl5//znP3FTLlq0CDcE4Icbt7i4GPjhblZW1l/+8hfc7kuXLkXCkf0hsrOzW7JkyWKCjh07hpqRxA/6E5KJmFetWoXGvz179mAeR0dHdHfmzJn29vY4vVOnTpwulJ1C4tes+CgGG+WypHKdC+wvKTXW2Lm9pXtHxx5GDr2MbYIMbYMN0Me+j4ldqLFLX2O/GFFAnDggwaZ7go1vrMirv5nHAJFfol2voWYp+aLyOb5FReYr54pO1lqf2aR3cl2XI2vFYHZ7V3vULOm6Yk7XaeXdZlYOTc5N9h8UP7xsxqo931548OwFTb+Wz0Fltij/EAGPolak/RUVFaEqwZgHkOD0/Px8XPOuXbvidJjm4PS4uDicXlpaysYvNDSUYyCSqyEWizHbrVu3OJuetL++fftiPcHWUdcB+sc//oEGS43E9p9qEa5y8IPU5tfSev9eXS3d9K26yvHraWgfZGDXy0hmhb1kc06nPsY+UebdI82CY8zDEsXhg61DkkRBSaLgZFFYYuecsVbz5zitmGe3bp7FoU02JzZ3ObnF+HCN64q5LtOnehcWeWeP6zV8TFhKbuSA0Qn+KUNtYgZbDUh1GZztk14cW1I1s/bgyRt3ZRDKfGmT0kGofMvUqVNJ/PCtBQsWYPvLzMykCPsj8cP8JSUlbDkM/NRGFsKzkfYH+Km1pJLI8S88PBynNzQ0wLwJ3frb3/726tUrmji9z0ftiFYgIZRi8KTSphM/Hp06qyQmKdzcsYup5FPAT9LT2KFHZ4fALoCirQxII5cQI/8oi8D+FiExFn1iLcIHiaKG2PUbbheZapOQZj5/tvu2Fa5fLbPftdJ5zyqHw+vNT24xPPql28oVoaOKgsfNSB5TET+0KCEyNyUka1iPjFHWAzNMB2YZx8JnlChlvPPIaW6Zpd2Gj6n+5vgrebDYhEY/YrsK+U/k93C7Q/1h5onbFGaVGvEDm2Pgx7A/cM7YpEiilP4T6YDtjwz4aPXxj8QPDO6DDz5g4Eepj3/sQttRNB5EaEq1Fa6YlzdJXxeXFVrYdzZz6GjupAeRn5W7HuBn66Pv2NNQEghWaATzFwDSL9IiKMYyNNaib7xVZKJ1VJJN/xS7QWmS5HTRxLGiQxs9T6yxPLRCtGel0941zkc2iE7Vindv7p4zKWTIhGGZk9PyKpKzpgxLKBkdkpfjOni4KHaE+aA884R8y4Q825QCt1HlDhllkmHFktjMklXVx6/fvvrbQ5ktKrcroGnY/hNVuLKysp1y0iGAH/ZIgB8eF7Eccv7C9p+00jdy+k9sRujnvn37SPywEMCPtD8cvdDsCQrpPxkRMdmboNyb964CeDDnlE1bnDtYuOqJPGT4WXl3tPLVgymMfZCxXaChXY9OTkFduvY28A8z6BVpGBJtHBlrFDvIICnRsCD9033LXE+vMj+zrMORpfr7FpvsWmSyf7n5t9W+JSVd04viUktGZE1Jy56SkF6anFSSETmhyGVoum3iKOukAtukAvfULwJGlQTmT/fMnmY7dII4Pts5Md07fnBibt4r2ZZGM1YY8MMzQ5iP4GrrZH9AyP4YdkzaX0hICMMUcCuT/vP27dtsc6Hl8xfcP2D8w9lhdsr2nwwJnPbHCKRUDreZath5YKuFpLOZsx7+WHbVt/DUF/t1tvTTs+rWCSC0C+gi9tNz7mXiGWzcPaRLcJhe376fJg38bGx6l8UTbA/MN7xd63h3i9mtLaZXq82ubDG7sFl0dJP3hGLv1LGxw0tGZZYNHz0leeTUIYNL02InZnXLmyAZmmOblGeblO8+bHzAqLLg3CnBY2YGj53fLadCkjTKbXBG1+QUmN3cefIYxkJcJWR/qGkg1MMVhlkobtP09HRt8MOBh8D4x8APWRin/2Q0vYD/fP/99xn4MYgDP1rtQAo5r6NfNbzoHtzV0qmTmQScpz5cWDjpmTp/BhBa+XYS+XQUebcX++jZ+HWx89P3DTYMDesSP8A4J81m2ljHVVOcasvFuyssLi//9OEWgwebO9/aZHx9k8n1jYbnNlhXlUnis0KjC/Pii8ZklKSNnjosY0ZuXGl+5Ph814wSh8Hj7ZILPdIndh9dFpRfEVI4p/e4BeHFC/pNmhswYqxj3BCH2GS3AbHfnzsHQY2A/0Q0a9YsnK7N/IUc/8COkRCB+SduVtAB2x9ASI5/ZLsz8MMSsP+EvAz82I6a8J8EYBhCOR/97NkTnwBXC4eOls4drZw62rjo2znrO7h9buupL/LSs/H63Nbrc2svfZG3gcT70+Q4o+n5NouKHJZNkqybbLut3HJvhfGBGSZ3N376qKbT461GdzcZXN1g/NNG+5VTrAsL/KPTYwYVjU8qKU4tLUif+kXmjEkpU0qCR40E2HxGlHZNK/YdMTEwpzy4YHqf8fMiJi6ImbwwYeriXjmlNgmjrRNHSgYNPXz+oiwoVFaJjP8wftAKXl5euE3xeMbGD7cIw3+S+KFEvviBUvefCD888mE92fghYs9f+EpRw4/rHK1U+UxY88FDO5w8TW1c2tu5t3f20HP37Ozjo9+1u5mdd2dnv47wsfM1tPIxdPF+f/5El9pSi+riLptLzGunWOyeBuAZfbfA+v62Lg+2Gj2sMXq01fjSBsv1M8TFxd3GTk0dOTEjq3RM+pSS1Kkz0qbPLly0cMaWDdOq1w+ZtSKyaFrvvLJeOZND86dHFs8dOKVqUEVVYsXipGnLvbPKREMmGCeNdU8tuAMz0eYG3ENJ+zM1NQ0KCurRo4e+vj4GDwL5n3/+mQ8/LKewsBCnQyzB8J8gp3PnzgMHDoyNjYVIES5GjBiB85L2FxERgdgQ5w8//EDaH+oKYWFhuNGfP38OYR/K/ve//x3CCbaXZntU5vN/FOGvZcYofT179mQ/f3tPT+Pkwd29ffUD/doH9Ojk7KPXK8Jt1pxxeSOj+/fSH5eo/1WF075y8c4ys+2TzXZWmB6YYXh0TuefV5jerzV6srX9862f3t9qdq62x87V0RuWp1QuSCubmbNw/aKFG1ZXbVy9bPumpV9Vr9m3Y++pk7tPndxw+PuJq6vjJk5PKK9MqlicUrEkpWJR8rQlMeWLuo4qc0orkgzKXLpzv8z05PghB8W3ft1OvhLm6el54MABsgnY+MHdO3fumJub43SIHRn4sQngpLjmLyQBWjU1NYiNEb/TyuFzxowZOP2zzz7jhINMRN8czz8ofyK5TVLq9U8/Hblw/vDkyVl+fhaBAeZBgVajcuPqHlyk6Mer5mcvLfbZNM7o8EzLQ9Otdk8W7SgX7Zsl/m6O6bF5+tfWmf+61ehJjd7Trfq3a8SrizvOHGc9rzxwzrSYSWWpU6sqFtZsnb25evLa1YVVC4rBBFevm7R4cd7chcNnLumeOS5t3prsRRtzF27MqVo3umpt/JQFgaMn9cgYV7p6y1NatuVEjt1k/M6gjz/+OD8/H7AhD2q6u7tjBnBcwPPJJ5/89a9/JTMeP34cZWHjh+eQaK0LsQF+nGvcQFu3bmX7T/CTH8sJZi5kxsDAQOEzpaT/1EhS+Z5O/d7d1cuWzFy1ev6+PZuk0ufy80gvzx1avm2q9/axn5+YYfTDHKvjc+33l1vsLTc5Wml+qkr/wRbzX2tMnmzp9Gt1l5/W2E7O6jB2tO2YkZKCbN/8wti88kkByWnug4Z2HzYyJCs3pqA4Z9bi/LmVo+YujRw/021oYd+x0/OXVBctry5bs3XKhu0Tlm8qW7np4NlLzym6kUZdTOVYSPv797//3b59+44dO7733nu4uQEehAfb/hiQIJJIJM3K+IT0n2DN0O5/VxL4arb9ARu6Cy4RmOF727ZtbP/JqQB8b9myhdNbcvtPYVbl9EkqX3hslK9+NFIyu4TWa2z49fTexUM2Fzl/M9vj8BxHgPDMAvGPVaIjc7pcXmnwqMbwwVaTR9Udfq02qJlimpvhNDgtKHlYt6TBXhHxwZ79B/oOGe03LKdnRn543oS+BaUDJ84A8JLK5wfmTHYdPt4rY2La7FWlq2umrts2Y/P2ksWrdx7+vommFDt/NM23/wDxA9rpfvHiRXZ2Nm4sNzc3jB/YH5+tAFlYWJw/fx4Dw4j/muQEMw74Bow57e/69euIAUYyuMCzUPb+EakGOICCggK+ZTM2UsznH4jrZppYf8Nn5tExMlpxChcuXlLPbzy5su/64bkHF0Uene9xar7NT1WWF1dY3l5v8rja4F6N6aPq9nfWd/pqtouvr55PeEBwTGDv/v5+UUGinn2s+sT7JI/omz0uYkxp4GiYeY4dMnN579zSntllHhkTfEZODh1TUbJGBl5l7a556zY9e4U2IhoV+7qU6owCuf9Hxg9PnjwBC0Dp4KbQJhwQuX6dlZVVVlYWGRmJUwAw0oOR9ofnn2T7oBT2/JMxMCH8sJ5isRjKhRnvRx99hFK8vb0Z/LSw/+S/xTxNKxMka7RG5WMZit0c2Q2YRzQ9oJ+dvPX9nFv7Rpxa6Hh+of6vG40f1pjcrzH9rbrThaUdDi72zkx1cw9w8Ah09w318gzvZdUzxLhbX5vQuMj8Sb2yJ/hkfOGSkuOdMb7HaBl+/iNLAkaXB+dMTa1YPKdmz3fnrtx//lJWZLOUGPZUx16w/UHroK1apDMAhr3ov/71L3Q2guaKH44ePYoj98mTJ2PJFCt+YL8PBH0z1s9orokF5/5DcHAwkg89DCIH7Q9cteD9rXjhXxZmyHZZ0f6g9JV8r/Ul3Xj2zJre15d3frLF8D6Mf5tN7m80urjSZP886zUzeoT1NggIcfaLCHAMDjT0CjDwCTXoFuGSkOkyeLR9XJbDwJEOCdndRk/pmVOBVl76jJkRVTi9dGX15QcPpWhHBJ+Rkaqpyoj/cHXAgwF+qOlhXMT4kf4TzT/r6+sx0miRGgsn8WOvn2ESiUQM/NjE3n8AtCZOnIgzQqSBhWuEQyB+oBkXmFPZbuT2dz0tO6RLy/c8Hl3alXVzndXDLQZ3q43vbTAER3p9vejbOUY1FY5jMyXB4c6uIb0MvHoa+IYbB/QX9UmyiEqxjB5mHZNpH5dtn5DrlFoUkDc9KHdqn/zysLEz+xXPHVY69/aTZ/KNeEoRqUqb5YtnKuX51q8Rfqi9wP7wuiJpf2fPnkWJLi4uKEVPTw9tl2P8MNik/ZEtQ66/tGPtP9BKZ8u5frZ7927cP2bPnk22NnnB/uZ9/oHYjsBLM2Q6hVpTPgrCdQN8mijZLiuY4OVvKy5u7nmr2uH6RsvbG4xurTe8uc74wkrLrZOMqiZ1C410M/HyNgqIEoUm2vQdLI5KsYoaYhU93CY2C4zPJjFPnDjGccg4r/Qi34wi78wJvlmTemV+cen+b7LZk7RB8a4ZxS6uSjHSfyL8UCsDDHgLns9/njlzBgmBYBwLuXz5MhbO8J98jSY8/iHiXH95/PgxTGtR4oABAzBC5AUn6X5+Xp0o5eEiCncWcKEvf6Yf19JXJv2y3g/Aq1tncHNtlxtrzA7PNpubJxqc6NYtqqdTWLikb5QkMs4hKtmu/3Cb/mn2gzIdEkdYJ+RaJRSI4sc4pxR6phe5ZExwSZ/kn1p46cFT+THfJvl5GMXeMtn/SftD+w+IsP1h/FA65/rZ2rVrMX6rVq2ilN2fc/+W2RSs/T9ONk78oKth3wsBCWP+SfF7UW3+/wP3w3mU+rXSzuXnG2R2eI9+fODs9uy69cbwubXO+NZ68x+qzLdNd5uU55uZHRafERkcF+LbL8I9KsEpJs0hdoRTfBZ8bOJzrZO+ECcW2ibkOaYUOGdMckoviRk/+7lMERl+zcRJX4rQXMB/AmwoHYDkww8xX716FfvJ4cOH43YU3n+geeaftPoElYEfuf8OdxMSElD6O++8c+fOHU5E2HEF8p/N5GyTVEh5fknKmosSsxj5ySKl8ckat0m+a09L7zXe+6Zug3ndepOba00vr7PbPctu1hi7gkyPjMzuKWk944eF9h8SE5wY7zUw2SEaHGmaRdQI8/4jLGOzRbF5VoOyrZPz7YaOc0kdn1+57rUCP2kTKhsFM1zzz3bq8QMAhu2PnL9w4gdgQ9SPEh0cHLAQTvzYROJXV1fHecKab/8Bn/MAwottFBGlkJvAKv9JgIcRYox5jECCptX48YEixV2YjjbJz2vT1G9Pb31zbZ2obi0MgeYnlrvOGus2Mt0/dVjP5KHdE4f0iEkKiBwUEBHbs/eAYKfQPpbhKUZ9RxtFDjOLSreIzjQfMMpk4CjL2FF+Q/KvPG6QL4TIZ+rK0RdfIyLxA9M5LSeYy4EvxSb1wQcfAEKIn1w/A/xw6wA8KPHdd9999uyZrvjhNZQ9e/aAAjCywvepU6egCBTps/FDRQMPTs/JyaG4DtizU8jnH7R/dlnFTKltXUhVk1LpC5qqu3h8xY21lrfXGgB+NZON8oaKhqZ4p6T4DB/injHMc0S638gM78KR3gWje8alhNr1DjMJG2oWmSqKHCrql2YRk2Uem20XOzKlaGq98pCHokgp07dQguvXGD9HR0eKf/8BtQg5lQfYKPX4vZ1y/4jdKBT/+jXqOmjpgDw/gfcfKPlWF94/8vT0JC1PgLjXr2meEZFChsUaCBXtSclPhiky1tPN1345OOvuehMY/G5uFB1bJN63yHv30pCvFgbtWth9R6V/7Tz/2rndN1V4LhzvVjiia0RcgHVPf4Me0cY9Y0yCYo36DDHvl+ESk7G8dm+zfDOLIg7wqh7fRb1FKoX4j7Eexv45a9YsXEe+/XfSPkpKSlqJH9bhww8/RFtCpHzAj2xwX19flA5+HoJRNgrsctspUeF0lRTLtVKy9RfVIVpasShKvExJvkJDyyKK5hsPz615sFbv5hqj6xvEZ5cbna7qcHrBp6fmf3Zi7idHZ33y7YyO2yfqbSjsvKrQdNyQLsnxduEDfCyCoky69TXs1k+v+wC97rEh6UXPmhX9BpWLHpxX6EK4EwH7aydfV0xLSwPnidch2fNP1A8ePXoEnhOl9+7dm73/oOX+LYM++ugj9vMP5P4tUG5uLna/x44dw10NM/D5z7YjZfvKI8JnRw99eWJ5+KU1PjfXWl5caXK40uzAPLuvZ9nvnelQXW6/qsh6xRfi5YXiykLXrESrxFjH2BiX8ChPSU9vg259unSPsQ6Kvv7oqfwUPfGCNKUTZXRJGD8qKythFlDJoo0bN964cQMNP1hTmCNght9++w3fgjaCyAGJgnACeTCYTGLhO3fu5LO/9evXI4FVVVX4G9GyZcuQAjdv3sSJu3btIrOfO3cOpUNBFy5cYAjnLLSN8VOsjshsU7aW9vL53RdXttJ3V97bEf3LKqcDcyz3zrLZWWFTM8l83Tjj1ePNlxdYLCu0mZMvSR1oOmigY79o19AIL/++gdYBvay6h4+YMKWe8bpRSv0vV9RD8mscP3hqocFrCfAwXlzE1+58YjkZSOEMHu71F1IEi5rJIZAzl/zRavkzs7Kdpte0tO7K16XfLem9Y4bj1ilWGyZYrB1nvmas7AP4Lcy3Lc9xSoyx7Bft3CfSq7B4dNn0svwJEzd8+VWDYnGHqYbqjU9cy7w0y+FwNhlmZj8yzvmObHZrkA9jkjzkN59u5IPBbLRI3UhV2WW14P1L7HiDTFeWpAgaZava8OfB1aP714+pGCFaOV68ushqxTibpWPtlxbYLsm3mpfnMD7TLSnOq3zqmO+Pf91EvUJbVDgyUX9rIaW+mKdqcc7mJvFg3BV+DlZAMiJGTCaAmTbEpwO+5lO1xf6Tacg40lfdkg+E6Ky7bN3k5c2+3duXpYsqc0yXF9osLrAHy6vMtZ6d61o1dci162eaZS63USptUsYhzWqvLGQ6G+bLfji15EvnIy35BdgoHnMnU7TXSiMnb/zAL1FAKDITlRXK93rk59yl9I266+F93Ib3N6/Mtlqda7i2wGh5vklVjvjLqiH0q+uK/XTZsym08g1qMnHNUuVSJ3MUlOGHgk1KPVRiWAOtqdUoYosArmEu8+uvvz58+BCmiw/kBFH8y5cv4QI/Vw08jx8/xhnx9jpECFgZNNd9+vQpSIO8aL+eVq7ywCwXCX/y5Amw4aeN0IPgWA5DZ3YVeMc/fksn4w0Fu2xzXmHazKU4RWghe7y9+dyF47199eaPtttSaLijWH97iWjNBO8X1/fATFVhppTijT5ot1++U8sujoOwwiQS2Nvw+U/yAhPM72/I6eeff758+fL169fv37///fffw8XZs2fPnDkDrb9nz54rV64cPHgQPaVw6NAhmC7C9HL37t3o2DzIrK2tBfC+/fZbEHXt2jWAZMWKFais06dP19XVgcwff/zx3r17kAgTXRTwAf+dO3dAIOofFNE1OSve1vGDjFhLOaBDM6Dz+uxP+4P8OmX177SpxOmrMrsvp3lWjg+i669AsKH0kU1sceq+kyVZO8JtwbhmswFC6Bpm88gOoPnQA81Ax48fB/wAWki8evUqGBBcQKyGNl0BmBMnTtByi4T4BGwLUnBbHzlyBC0UoO0qiEnAjuEn9A8QhYr47rvvoB9AES9evMAqCdRLt/e3qsviWHLjLEz+GomXNVsWh4c5BQUYhvZo/0WG08g443Cf94P9jZoa5LEXIzBQyZeq/rmE0m0SiGoID4Qrz8m/fft2QA7M6JdffoGGhm/wn4cPH4aLHTt2gAlC4wIPALZhwwa4vnTpEjhbQBRaHNA6f/488MPFxYsXwf727dsHKQAPwAyJwAyQM/CDPgGYgUDgAUcNRgkdAp/TEa6OZv+Jc6ozqDaVGJyMvJTs9SAvJ0wZ4+Jh5Oj8uZ3kY0vbf4vs/uXm2dHb2zgs3K+xqV7Bi06cqr00SN0bq790lFKuYgvrz2d2DE5a6UjxXi4AhuwPGheb0cmTJ+EnoIhWasCLfvPNN8AJ+AEAyIbQq0XgFgCJUlCkAULg++jRo5BIy9cEAD9I2blzJ2QHnwmeE+1awEALEjBgAvVqcfxAXgu8NJBuputnVZUb23Ywsv0YPgbWHxvZfGZo84m9h0FwhPfZcyeVaqnkKOMFxRE3SvWyc8VYSP5/Cc7+x6EHDxuZjoYZZCvwE8YtGKKgHa/KCfGAAYHNwS2ENwx76B0HQAAYYgMDAsuDaQgYExgueq0MDG8wCkKuV69egSOl5JMgYIPxD76RGiAKioPswA95qTfx/laWy+J4bydxTe89uMfWRWRk09HA7jND+8+72HQwtOkYERcweFjfg99tA+tU8ioDDdSgalGDcluDw8eqLIl9Jozdc4XBRnfxmV3lMTuVQPInTiHLYjQ3WwKpMMlPZic52fE7g9T8J9n6fBm4ZDHxo1SvBqVv3rrx7eH9Enexqb2eoe1nhnYdDK07zFxQVrVoemPTcynVoABEqtbESjnoWqo4babYbVAUR+rM1oohSqcacbYXZz9QawX+WSJZOl61IfUUKF34WuD8Ei+xjFpgy4KWv6yw+dbta2kjk01tPwf/GREb9Mulszdv1SnOPtF4SwGHCqqVF1UTKJdgiEGXO7QQwI/RoOx0PiLZ2EbGTsT8DFAFMCbvMsxOQLHWxw+cM0D1B3rlTd/U/Kpq2Swjq/YpaXFPnj1UmZNaNoqQwFec+nZEWxPFEzJTXOarEzN5S4BBowSSdDu/y7qrmn+y05VXzcqnCel6aX1Y/5B5lTNkw57sMJkKDNV8BRVEhH2UmtNQvsiOX02y22pZL03V1IGfjSgfxm1CKvwEjFSAyIw00bkYPAinx48f+vl7Xbl6QXkgSrUhR2bk1EdLJdk60OpdnqE2+4JdFp8+fD/Ja8YMRXsSKJG81Wb/f5pzpksUKd+MkDY9evQbGhHV19h0foO4AL/wGNOazqoN6VoR7YlT7TexfsZHCptT7TAQRLXOq/y+2d+QKG2I+/kHASU0cjK6DE7EW4LarNoIpDDK4mRjqCdcKXZednEkD2dxfFnInxqLY9eR3RqM6rQkfmC0mkA6Z6mYOLe5heUziM9ZMZpPo2KcEjTy6KSqrsycpbNjlbfpP4WIEuyS/0X0livy++DHrqSu1Rbgf0MtqKXY3wE/XZ2AsEPgTCGrx3lXmVEtouBk00YfPq1ollNiX7DL4tOH7ydfNTUqL8zPqWo7Pu7WEGdJwiT36W3w/480Tt/btqa6ltWClhGmtvefFNHZdSSdnsFQldXS4lRCcJuyJQukCKv0dkht/aVNJHI2hxZsUnZQT3G5NUZ2vuLIbiSsjLC2ZBdhKCbAz9ZNWFuNxFcR3v8/3SakjWRCAa38p0aFORnYlddSwzfaPq2n/8UPbUw6VaT1tf6vxA91vTenDKMsxoWW/G+HWhs/kCQghKwe5y2GEFrdvwlk5xTF5uSTQHM5UrZKAhnZiQL86Jtzy1ejYpwlqvYf3sTCuTbdovVdR/jhgTYvrjUkDEYL6E3FD285V8uyt7L0Ny1KG+LYv22lREZ308iGf9KsyjO6Lc3lndBPxnNiDH5tNBG4S3PZijA/zWoH9oVOxFeRP0T8oKtMjVnYDLjyjE2PN6ShTvytoT/K/LMNiWpTD9a20tqc/ij46dpMAvxvCD8txWpkwzbaJtTa80sksUVpI5asNqXu5TnZNErjLJfdfNqL1bV9BPi1l8OpKputzc4vkaTToSstqyTciHzHZ1tcXGuyaFN6mxgM9f/M+Ee1dEbX4uxtqElriPf9rXyEOYUZ2NLIdPZPhgRavXuSiZylsEWxc/GpzZdXoFKcmjPKYuvPuOAjzo7FVx0O/NqQOGHQhodPE4EjT6QEvhbkS9e4lEVm1GmLWOAZbi2FCFf2j+I/KcEu+V9Eb7kiLXh/eauITz6l7ig0CsF9U/siWkY6KdbmpWsktf/foZOB81k3rd6+jPpwSqDVnRsnZ8vU06YsTrUFfnLqQ/Loqipfs2gjvw3iBwHPrk1N2gSYPw69uecfOOntjX+U1o6F4jHclknTJkvrhbdJ3haIfXvxA0MPko1xl+LxHnzZtRElUDW+vHxsnMUxcjGshJ2dUxNOOcLV+T9mdcbhBsgIFAAAAABJRU5ErkJggg=="


def gerar_pdf_proposta(dados_front, dados_orcamento, contato, orcamento_id):

    cliente = dados_front.get("cliente") or {}
    endereco = dados_front.get("endereco") or {}

    def primeiro_valor(*valores):
        for valor in valores:
            if valor not in (None, ""):
                return valor
        return ""

    def float_seguro(valor, padrao=0.0):
        try:
            return float(valor or 0)
        except (TypeError, ValueError):
            return padrao

    def numero_pt(valor):
        return f"{float_seguro(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    def texto(valor):
        return str(valor or "").strip()

    def escape_html(valor):
        valor = unescape(texto(valor))
        return (
            valor
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

    def endereco_formatado():
        partes = []
        logradouro = texto(endereco.get("logradouro"))
        numero = texto(endereco.get("numero"))
        complemento = texto(endereco.get("complemento"))
        bairro = texto(endereco.get("bairro"))
        cidade = texto(endereco.get("cidade"))
        uf = texto(endereco.get("uf"))
        cep = texto(endereco.get("cep"))

        if logradouro:
            linha = logradouro
            if numero:
                linha += f", Nº {numero}"
            if complemento:
                linha += f", {complemento}"
            if bairro:
                linha += f", {bairro}"
            partes.append(linha)

        cidade_linha = cidade
        if cep:
            cidade_linha += f" - {cep}" if cidade_linha else cep
        if uf:
            cidade_linha += f", {uf}" if cidade_linha else uf
        if cidade_linha:
            partes.append(cidade_linha)

        return partes

    styles = getSampleStyleSheet()

    normal = ParagraphStyle(
        "PropostaNormal",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.8,
        leading=10.2,
        spaceBefore=0,
        spaceAfter=0,
        leftIndent=0,
        rightIndent=0,
        firstLineIndent=0,
        textColor=colors.black,
    )
    title_style = ParagraphStyle(
        "PropostaTitle",
        parent=normal,
        fontName="Helvetica-Bold",
        fontSize=15.2,
        leading=18,
        alignment=TA_CENTER,
    )
    table_header = ParagraphStyle(
        "PropostaTableHeader",
        parent=normal,
        fontName="Helvetica-Bold",
        fontSize=8.2,
        leading=9.2,
        alignment=TA_CENTER,
    )
    table_left = ParagraphStyle(
        "PropostaTableLeft",
        parent=normal,
        fontSize=8.1,
        leading=9.0,
        leftIndent=0,
        rightIndent=0,
        firstLineIndent=0,
    )
    table_center = ParagraphStyle(
        "PropostaTableCenter",
        parent=table_left,
        alignment=TA_CENTER,
    )
    table_right = ParagraphStyle(
        "PropostaTableRight",
        parent=table_left,
        alignment=TA_RIGHT,
    )
    empresa_style = ParagraphStyle(
        "EmpresaHeader",
        parent=normal,
        fontSize=8.5,
        leading=9.5,
        alignment=TA_RIGHT,
    )

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=8 * mm,
        bottomMargin=18 * mm,
        title=f"Proposta Comercial {orcamento_id}",
        author="BRFER Comércio de Ferramentas LTDA",
    )

    story = []

    logo_stream = BytesIO(base64.b64decode(BRFER_LOGO_BASE64))
    logo = Image(logo_stream, width=38 * mm, height=28 * mm)

    empresa_html = (
        "<b>BRFER COMÉRCIO DE FERRAMENTAS LTDA</b><br/>"
        "40.954.410/0001-96<br/>"
        "www.brfer.com.br<br/>"
        "(11) 4362-5151<br/>"
        "Rua Coronel Francisco Rodrigues Seckler, 53, galpão<br/>"
        "Paulicéia, São Bernardo do Campo - SP<br/>"
        "09.693-050<br/>"
        "799387168111"
    )

    header = Table(
        [[logo, Paragraph(empresa_html, empresa_style)]],
        colWidths=[45 * mm, 145 * mm],
        rowHeights=[28 * mm],
    )
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(header)
    story.append(Spacer(1, 6 * mm))

    numero_proposta = primeiro_valor(
        dados_orcamento.get("numeroProposta"),
        dados_orcamento.get("numero"),
        orcamento_id,
    )
    story.append(Paragraph(
        f"Proposta Comercial Nº {escape_html(numero_proposta)}",
        title_style,
    ))
    story.append(Spacer(1, 5 * mm))

    nome_cliente = texto(
        primeiro_valor(
            cliente.get("razao_social"),
            cliente.get("nome"),
            contato.get("nome") if isinstance(contato, dict) else "",
            "Cliente",
        )
    )
    documento = texto(primeiro_valor(
        cliente.get("cpf_cnpj"),
        contato.get("cpfCnpj") if isinstance(contato, dict) else "",
    ))
    endereco_linhas = endereco_formatado()

    telefone = texto(primeiro_valor(
        contato.get("telefone") if isinstance(contato, dict) else "",
        cliente.get("telefone"),
    ))
    celular = texto(primeiro_valor(
        contato.get("celular") if isinstance(contato, dict) else "",
        contato.get("telefoneCelular") if isinstance(contato, dict) else "",
        cliente.get("celular"),
    ))
    email = texto(primeiro_valor(
        contato.get("email") if isinstance(contato, dict) else "",
        cliente.get("email"),
    ))

    story.append(Paragraph("Para", normal))
    story.append(Paragraph(escape_html(nome_cliente), normal))
    story.append(Spacer(1, 3.5 * mm))

    dados_endereco = [
        [Paragraph("<b>Endereço do Cliente</b>", normal)],
        [Paragraph(escape_html(documento), normal)],
    ]
    for linha in endereco_linhas:
        dados_endereco.append([Paragraph(escape_html(linha), normal)])

    contato_linha = []
    if telefone:
        auxTelefone = escape_html(telefone)
        auxTelefone = "".join([char for char in auxTelefone if char.isdigit()])
        if(len(auxTelefone) == 10):
            auxTelefone = "(" + auxTelefone[0:2] +") " + auxTelefone[2:6] + "-" + auxTelefone[6:]
        else:
            auxTelefone = "(" + auxTelefone[0:2] +") " + auxTelefone[2:7] + "-" + auxTelefone[7:]
        contato_linha.append(f"Fone: {auxTelefone}")
    if celular and celular != telefone:
        contato_linha.append(f"Celular: {escape_html(celular)}")
    if email:
        contato_linha.append(f"E-mail: {escape_html(email)}")
    if contato_linha:
        dados_endereco.append([Paragraph(", ".join(contato_linha), normal)])

    endereco_box = Table(dados_endereco, colWidths=[190 * mm])
    endereco_box.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.75, colors.black),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0.4),
    ]))
    story.append(endereco_box)
    story.append(Spacer(1, 4.5 * mm))

    introducao = dados_front.get("introducao") or (
        "Prezado cliente, seguem abaixo proposta comercial com "
        "pagamento à vista com desconto e nossos dados bancários:\n\n"
        "Segue nossos dados bancários:\n"
        "BRFER Comércio de Ferramentas LTDA\n"
        "CNPJ 40.954.410/0001-96\n"
        "Banco: 341 – Itaú\n"
        "Agência: 8811\n"
        "Conta Corrente: 99874-2\n\n"
        "Se preferir, o pagamento pode ser realizado via PIX, a chave "
        "é o nosso CNPJ \n"
    )

    for bloco in str(unescape(texto(introducao))).split("\n\n"):
        linhas = str(bloco).split("\n")
        html_bloco = "<br/>".join(escape_html(linha) for linha in linhas)
        story.append(Paragraph(html_bloco, normal))
        story.append(Spacer(1, 1 * mm))

    story.append(Paragraph("<b>Itens de produto ou serviço</b>", normal))
    story.append(Spacer(1, 1 * mm))

    itens = dados_orcamento.get("itens")
    if not isinstance(itens, list) or not itens:
        itens = []
        for item in dados_front.get("carrinho") or []:
            itens.append({
                "produto": {
                    "descricao": item.get("descricao") or item.get("nome"),
                    "sku": item.get("sku"),
                },
                "quantidade": item.get("quantidade", 1),
                "valorUnitario": item.get("preco_unitario", 0),
                "descrComplementarOrc": item.get("descricao") or item.get("nome"),
            })

    corpo_itens = [[
        Paragraph("Nº", table_header),
        Paragraph("Imagem", table_header),
        Paragraph("Item", table_header),
        Paragraph("SKU", table_header),
        Paragraph("Qtd", table_header),
        Paragraph("Un", table_header),
        Paragraph("Preço un", table_header),
        Paragraph("Total", table_header),
    ]]

    soma_quantidades = 0.0
    total_itens_calculado = 0.0

    for indice, item in enumerate(itens, start=1):
        produto = item.get("produto") or {}
        descricao = texto(primeiro_valor(
            produto.get("descricao"),
            item.get("descricao"),
            item.get("nome"),
            "Produto",
        ))
        sku = texto(primeiro_valor(produto.get("sku"), item.get("sku")))
        quantidade = float_seguro(item.get("quantidade"), 1.0)
        valor_unitario = float_seguro(
            item.get("valorUnitario"),
            float_seguro(item.get("preco_unitario")),
        )
        total_item = quantidade * valor_unitario
        soma_quantidades += quantidade
        total_itens_calculado += total_item

        complemento = texto(item.get("descrComplementarOrc"))
        descricao_html = f"<b>{escape_html(descricao)}</b>"
        if complemento and unescape(complemento).strip() != unescape(descricao).strip():
            complemento_formatado = html_tiny_para_reportlab(complemento)
            descricao_html += f"<br/>{complemento_formatado}"

        produto_id = primeiro_valor(
            produto.get("id"),
            produto.get("idProduto"),
            item.get("produto_id"),
            item.get("idProduto"),
        )

        imagem_pdf = None
        if produto_id:
            try:
                imagem_url = ObterImagemDoProduto(
                    produto_id
                )
                imagem_pdf = carregar_imagem_produto(
                    imagem_url
                )
            except TinyAPIError as e:
                # A ausência/erro da imagem não deve impedir a emissão do PDF.
                print(
                    "Não foi possível obter a imagem do produto",
                    produto_id,
                    ":",
                    e
                )

        corpo_itens.append([
            Paragraph(str(indice), table_center),
            imagem_pdf if imagem_pdf else "",
            Paragraph(descricao_html, table_left),
            Paragraph(escape_html(sku), table_left),
            Paragraph(numero_pt(quantidade), table_center),
            Paragraph("UN", table_center),
            Paragraph(numero_pt(valor_unitario), table_right),
            Paragraph(numero_pt(total_item), table_right),
        ])

    subtotal = float_seguro(primeiro_valor(
        dados_orcamento.get("valorSubtotal"),
        dados_orcamento.get("valorTotal"),
        total_itens_calculado,
    ))
    total_proposta = float_seguro(primeiro_valor(
        dados_orcamento.get("valorTotal"),
        subtotal,
    ))

    corpo_itens.append([
        Paragraph(
            f"<b>Número de itens: {len(itens)}</b><br/>"
            f"<b>Soma das quantidades:</b> {numero_pt(soma_quantidades)}",
            normal,
        ),
        "",
        "",
        "",
        "",
        "",
        Paragraph("<b>Total dos itens</b>", table_right),
        Paragraph(numero_pt(subtotal), table_right),
    ])

    tabela_itens = Table(
        corpo_itens,
        colWidths=[
            8 * mm,       # Nº
            20 * mm,      # Imagem
            76 * mm,      # Item
            20 * mm,      # SKU
            13 * mm,      # Qtd
            12 * mm,      # Un
            20.5 * mm,    # Preço un
            20.5 * mm,    # Total
        ],
        splitByRow=1,
        splitInRow=1,
    )
    tabela_itens.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDEDED")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2.2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
        ("SPAN", (0, -1), (5, -1)),
        ("ALIGN", (6, -1), (7, -1), "RIGHT"),
    ]))
    story.append(tabela_itens)
    story.append(Spacer(1, 4.5 * mm))

    data_raw = primeiro_valor(
        dados_orcamento.get("data"),
        datetime.now(ZoneInfo("America/Sao_Paulo")).date().isoformat(),
    )
    try:
        data_fmt = datetime.fromisoformat(str(data_raw)[:10]).strftime("%d/%m/%Y")
    except Exception:
        data_fmt = str(data_raw)

    resumo_data = Table([
        [
            Paragraph("<b>Data</b>", normal),
            Paragraph("<b>Total dos itens</b>", table_right),
            Paragraph("<b>Total da proposta</b>", table_right),
        ],
        [
            Paragraph(escape_html(data_fmt), normal),
            Paragraph(numero_pt(subtotal), table_right),
            Paragraph(numero_pt(total_proposta), table_right),
        ],
    ], colWidths=[47 * mm, 71.5 * mm, 71.5 * mm])
    resumo_data.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDEDED")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(resumo_data)
    story.append(Spacer(1, 4 * mm))

    observacoes = dados_front.get("observacoes") or (
        "Somos um E-COMMERCE, não reservamos estoque antes da aprovação do pagamento."
    )
    resumo = dados_front.get("resumo_carrinho") or {}
    total_carrinho = float_seguro(primeiro_valor(resumo.get("total"), subtotal))
    valor_avista = float_seguro(resumo.get("avista"))
    valor_3x = float_seguro(resumo.get("parcela_3x"))
    valor_12x = float_seguro(resumo.get("parcela_12x"))

    if not valor_avista:
        valor_avista = total_carrinho * 0.98

    pagamentos_html = (
        "<h1>Condições de pagamento:</h1><br/>"
        f"<h2>Total: <b> R$ {numero_pt(total_carrinho)} </b></h2><br/>"
        f"<h2>Pagamento à vista com desconto: <b> R$ {numero_pt(valor_avista)}</b></h2><br/>"
        f"<h2><b>3x de R$ {numero_pt(valor_3x)}</b> sem juros</h2><br/>"
        f"<h2><b>12x de R$ {numero_pt(valor_12x)}</b> com juros no cartão. </h2><br/><br/>"
        "*Frete a combinar. Entre em contato com nosso time de vendas para obter uma cotação."
    )
    observacoes_html = escape_html(observacoes).replace("\n", "<br/>")

    observacoes_box = Table(
        [[Paragraph(observacoes_html + "<br/><br/>" + pagamentos_html, normal)]],
        colWidths=[190 * mm],
    )
    observacoes_box.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.75, colors.black),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))

    story.append(Paragraph("<b>Observações</b>", normal))
    story.append(Spacer(1, 1.2 * mm))
    story.append(observacoes_box)
    story.append(Spacer(1, 4.5 * mm))

    story.append(Paragraph("Atenciosamente,", normal))
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph("Departamento de Vendas.", normal))

    doc.build(story)
    buffer.seek(0)
    return buffer


CORS(
    app,
    resources={
        r"/api/*": {
            "origins": "*"
        }
    }
)

TINY_API_URL = "https://api.tiny.com.br/public-api/v3"

TINY_AUTH_URL = (
    "https://accounts.tiny.com.br/"
    "realms/tiny/protocol/openid-connect/auth"
)

TINY_TOKEN_URL = (
    "https://accounts.tiny.com.br/"
    "realms/tiny/protocol/openid-connect/token"
)

TINY_CLIENT_ID = os.environ.get(
    "TINY_CLIENT_ID"
)

TINY_CLIENT_SECRET = os.environ.get(
    "TINY_CLIENT_SECRET"
)

TINY_REDIRECT_URI = os.environ.get(
    "TINY_REDIRECT_URI"
)


REDIS_URL = (
    os.environ.get("UPSTASH_REDIS_REST_URL")
    or
    os.environ.get("KV_REST_API_URL")
)

REDIS_TOKEN = (
    os.environ.get("UPSTASH_REDIS_REST_TOKEN")
    or
    os.environ.get("KV_REST_API_TOKEN")
)

TINY_TOKEN_KEY = "tiny:oauth:tokens"

OAUTH_STATE_KEY = "tiny:oauth:state"
TINY_REFRESH_LOCK_KEY = "tiny:oauth:refresh-lock"
TINY_REFRESH_LOCK_TTL = 60
TINY_REFRESH_WAIT_SECONDS = 12

class TinyAPIError(Exception):

    def __init__(
        self,
        mensagem,
        status=None,
        resposta=None
    ):

        super().__init__(mensagem)

        self.mensagem = mensagem
        self.status = status
        self.resposta = resposta

def redis_disponivel():

    return bool(
        REDIS_URL
        and REDIS_TOKEN
    )


def redis_request(
    comando,
    *argumentos
):

    if not redis_disponivel():

        raise RuntimeError(
            "Upstash Redis não está configurado."
        )


    url = REDIS_URL.rstrip("/") + "/"


    payload = [
        comando,
        *argumentos
    ]


    response = requests.post(
        url,
        headers={
            "Authorization":
                f"Bearer {REDIS_TOKEN}",

            "Content-Type":
                "application/json"
        },
        json=payload,
        timeout=10
    )


    if not response.ok:

        raise RuntimeError(
            "Erro ao acessar Upstash Redis: "
            f"HTTP {response.status_code} "
            f"{response.text}"
        )


    dados = response.json()

    return dados.get("result")


def redis_get(chave):

    return redis_request(
        "GET",
        chave
    )


def redis_set(
    chave,
    valor,
    expiracao=None
):

    if expiracao:

        return redis_request(
            "SET",
            chave,
            valor,
            "EX",
            str(expiracao)
        )


    return redis_request(
        "SET",
        chave,
        valor
    )


def redis_delete(chave):

    return redis_request(
        "DEL",
        chave
    )

def carregar_tokens():

    valor = redis_get(
        TINY_TOKEN_KEY
    )


    if not valor:

        return None


    try:

        return json.loads(
            valor
        )

    except Exception:

        print(
            "ERRO: tokens armazenados no Redis "
            "não são um JSON válido."
        )

        return None


def salvar_tokens(
    access_token,
    refresh_token,
    expires_in
):
    agora = int(time.time())
    expires_in = int(expires_in or 3600)
    expires_at = agora + expires_in - 60

    dados = {
        "access_token": access_token,
        "refresh_token": refresh_token,
        "expires_at": expires_at,
        "updated_at": agora
    }

    print("[OAUTH DEBUG] Preparando salvamento dos tokens")
    print("[OAUTH DEBUG] expires_in recebido:", expires_in)
    print("[OAUTH DEBUG] expires_at:", expires_at)
    print("[OAUTH DEBUG] expires_at local:", formatar_timestamp(expires_at))
    print("[OAUTH DEBUG] access_token:", resumo_token(access_token))
    print("[OAUTH DEBUG] refresh_token:", resumo_token(refresh_token))

    redis_set(
        TINY_TOKEN_KEY,
        json.dumps(dados)
    )

    print("[OAUTH DEBUG] Tokens OAuth salvos no Upstash Redis.")

    tokens_salvos = carregar_tokens()
    if tokens_salvos:
        log_resumo_tokens(
            "CONFIRMAÇÃO APÓS SALVAR NO REDIS",
            tokens_salvos
        )
    else:
        print("[OAUTH DEBUG] ATENÇÃO: não foi possível reler os tokens após salvar.")

def headers_tiny(
    access_token
):

    return {

        "Authorization":
            f"Bearer {access_token}",

        "Content-Type":
            "application/json",

        "Accept":
            "application/json"
    }

def resposta_json(
    response
):

    try:

        return response.json()

    except Exception:

        return response.text
    

def fingerprint_token(token):
    """Identificador seguro do token para os logs, sem expor o valor completo."""
    if not token:
        return None
    return sha256(str(token).encode("utf-8")).hexdigest()[:12]


def resumo_token(token):
    """Resumo seguro de um token para diagnóstico."""
    if not token:
        return {
            "presente": False,
            "tamanho": 0,
            "inicio": None,
            "fingerprint": None,
        }

    token = str(token)
    return {
        "presente": True,
        "tamanho": len(token),
        "inicio": token[:8] + "...",
        "fingerprint": fingerprint_token(token),
    }


def formatar_timestamp(timestamp):
    if not timestamp:
        return None
    try:
        return datetime.fromtimestamp(
            int(timestamp),
            tz=ZoneInfo("America/Sao_Paulo")
        ).strftime("%Y-%m-%d %H:%M:%S %Z")
    except Exception:
        return str(timestamp)


def log_resumo_tokens(rotulo, tokens):
    tokens = tokens or {}
    agora = int(time.time())
    expires_at = int(tokens.get("expires_at") or 0)

    print("")
    print("=" * 80)
    print(f"[OAUTH DEBUG] {rotulo}")
    print("=" * 80)
    print("Horário atual:", formatar_timestamp(agora))
    print("Access token:", resumo_token(tokens.get("access_token")))
    print("Refresh token:", resumo_token(tokens.get("refresh_token")))
    print("Atualizado em:", formatar_timestamp(tokens.get("updated_at")))
    print("Expira em:", formatar_timestamp(expires_at))
    print(
        "Segundos restantes:",
        expires_at - agora if expires_at else None
    )
    print("Chaves armazenadas:", list(tokens.keys()))
    print("=" * 80)
    print("")


@app.route(
    "/api/oauth/autorizar",
    methods=["GET"]
)
def oauth_autorizar():

    if not TINY_CLIENT_ID:

        return jsonify({

            "erro":
                "TINY_CLIENT_ID não configurado na Vercel."

        }), 500


    if not TINY_CLIENT_SECRET:

        return jsonify({

            "erro":
                "TINY_CLIENT_SECRET não configurado na Vercel."

        }), 500


    if not TINY_REDIRECT_URI:

        return jsonify({

            "erro":
                "TINY_REDIRECT_URI não configurado na Vercel."

        }), 500


    if not redis_disponivel():

        return jsonify({

            "erro":
                "Upstash Redis não configurado.",

            "REDIS_URL":
                bool(REDIS_URL),

            "REDIS_TOKEN":
                bool(REDIS_TOKEN)

        }), 500

    state = secrets.token_urlsafe(
        32
    )


    redis_set(
        OAUTH_STATE_KEY,
        state,
        600
    )


    parametros = {

        "client_id":
            TINY_CLIENT_ID,

        "redirect_uri":
            TINY_REDIRECT_URI,

        "response_type":
            "code",

        "state":
            state
    }


    url = (
        TINY_AUTH_URL
        + "?"
        + urlencode(parametros)
    )


    print(
        "Iniciando autorização OAuth Tiny."
    )


    return redirect(
        url
    )

class HTMLParaTexto(HTMLParser):
    def __init__(self):
        super().__init__()
        self.partes = []
        self.tags_bloco = {
            "p",
            "div",
            "br",
            "h1",
            "h2",
            "h3",
            "h4",
            "h5",
            "h6",
            "li",
        }

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()

        if tag == "br":
            self.partes.append("\n")
        elif tag in self.tags_bloco:
            self.partes.append("\n")

    def handle_endtag(self, tag):
        tag = tag.lower()

        if tag in self.tags_bloco and tag != "br":
            self.partes.append("\n")

    def handle_data(self, data):
        self.partes.append(data)


def html_para_texto(valor):
    if not valor:
        return ""

    parser = HTMLParaTexto()
    parser.feed(str(valor))
    parser.close()

    texto = "".join(parser.partes)

    texto = unescape(texto)

    linhas = []
    for linha in texto.splitlines():
        linha = " ".join(linha.split())
        if linha:
            linhas.append(linha)

    return "\n".join(linhas).strip()

@app.route(
    "/api/oauth/callback",
    methods=["GET"]
)
def oauth_callback():

    codigo = request.args.get(
        "code"
    )

    state = request.args.get(
        "state"
    )

    erro = request.args.get(
        "error"
    )


    if erro:

        return jsonify({

            "erro":
                "O Tiny recusou a autorização.",

            "detalhes":
                erro,

            "descricao":
                request.args.get(
                    "error_description"
                )

        }), 400


    if not codigo:

        return jsonify({

            "erro":
                "Código de autorização não recebido."

        }), 400


    if not state:

        return jsonify({

            "erro":
                "State OAuth não recebido."

        }), 400

    state_salvo = redis_get(
        OAUTH_STATE_KEY
    )


    if not state_salvo:

        return jsonify({

            "erro":
                "State OAuth expirado ou inexistente.",

            "orientacao":
                "Acesse /api/oauth/autorizar novamente."

        }), 400


    if not secrets.compare_digest(
        str(state_salvo),
        str(state)
    ):

        return jsonify({

            "erro":
                "State OAuth inválido."

        }), 400

    redis_delete(
        OAUTH_STATE_KEY
    )

    try:

        response = requests.post(

            TINY_TOKEN_URL,

            data={

                "grant_type":
                    "authorization_code",

                "client_id":
                    TINY_CLIENT_ID,

                "client_secret":
                    TINY_CLIENT_SECRET,

                "redirect_uri":
                    TINY_REDIRECT_URI,

                "code":
                    codigo
            },

            headers={

                "Accept":
                    "application/json",

                "Content-Type":
                    "application/x-www-form-urlencoded"
            },

            timeout=30
        )


        dados = resposta_json(
            response
        )


        print(
            "OAuth Tiny HTTP:",
            response.status_code
        )


        if not response.ok:

            print(
                "Resposta OAuth:",
                dados
            )


            return jsonify({

                "erro":
                    "Tiny recusou a troca do código.",

                "status_tiny":
                    response.status_code,

                "resposta_tiny":
                    dados

            }), 502


        access_token = dados.get(
            "access_token"
        )

        refresh_token = dados.get(
            "refresh_token"
        )

        expires_in = dados.get(
            "expires_in",
            3600
        )


        if not access_token:

            return jsonify({

                "erro":
                    "Tiny não retornou access_token.",

                "resposta_tiny":
                    dados

            }), 502


        if not refresh_token:

            return jsonify({

                "erro":
                    "Tiny não retornou refresh_token.",

                "resposta_tiny":
                    dados

            }), 502
        
        salvar_tokens(

            access_token,

            refresh_token,

            expires_in
        )


        return jsonify({

            "sucesso":
                True,

            "mensagem":
                (
                    "Aplicação autorizada com sucesso. "
                    "Os tokens foram armazenados "
                    "automaticamente no Upstash Redis."
                ),

            "expira_em_segundos":
                expires_in,

            "proximo_passo":
                (
                    "A integração já pode utilizar "
                    "/api/gerar-proposta."
                )

        }), 200


    except requests.RequestException as e:

        return jsonify({

            "erro":
                "Erro de comunicação com o OAuth do Tiny.",

            "detalhes":
                str(e)

        }), 502

def renovar_access_token(
    tokens
):
    tokens = tokens or {}
    refresh_token = tokens.get("refresh_token")

    if not refresh_token:
        print("[OAUTH DEBUG] refresh_token não encontrado.")
        raise TinyAPIError(
            "Refresh token não encontrado. Autorize novamente a aplicação no Tiny.",
            401,
            {"autorizacao": "/api/oauth/autorizar"}
        )

    log_resumo_tokens("ANTES DA RENOVAÇÃO AUTOMÁTICA", tokens)

    lock_token = secrets.token_urlsafe(32)
    lock_adquirido = False

    try:
        for tentativa in range(int(TINY_REFRESH_WAIT_SECONDS * 2)):
            resultado_lock = redis_request(
                "SET",
                TINY_REFRESH_LOCK_KEY,
                lock_token,
                "NX",
                "EX",
                str(TINY_REFRESH_LOCK_TTL)
            )

            if resultado_lock == "OK":
                lock_adquirido = True
                print("[OAUTH DEBUG] Lock de renovação adquirido. Tentativa:", tentativa + 1)
                break

            time.sleep(0.5)
            tokens_atualizados = carregar_tokens()
            if not tokens_atualizados:
                continue

            novo_access = tokens_atualizados.get("access_token")
            novo_expires_at = int(tokens_atualizados.get("expires_at", 0))

            if (
                novo_access
                and novo_access != tokens.get("access_token")
                and int(time.time()) < novo_expires_at
            ):
                print("[OAUTH DEBUG] Outra execução já renovou o OAuth.")
                log_resumo_tokens("TOKENS RENOVADOS POR OUTRA EXECUÇÃO", tokens_atualizados)
                return {
                    "access_token": novo_access,
                    "refresh_token": tokens_atualizados.get("refresh_token"),
                    "expires_in": max(1, novo_expires_at - int(time.time()))
                }

        if not lock_adquirido:
            tokens_atualizados = carregar_tokens()
            if tokens_atualizados:
                novo_access = tokens_atualizados.get("access_token")
                novo_expires_at = int(tokens_atualizados.get("expires_at", 0))
                if (
                    novo_access
                    and novo_access != tokens.get("access_token")
                    and int(time.time()) < novo_expires_at
                ):
                    return {
                        "access_token": novo_access,
                        "refresh_token": tokens_atualizados.get("refresh_token"),
                        "expires_in": max(1, novo_expires_at - int(time.time()))
                    }

            raise TinyAPIError(
                "Outra requisição está renovando o acesso ao Tiny. Tente novamente em alguns segundos.",
                503,
                {"motivo": "lock_de_renovacao_oauth"}
            )

        tokens_atuais = carregar_tokens() or {}
        access_atual = tokens_atuais.get("access_token")
        expires_at_atual = int(tokens_atuais.get("expires_at", 0))
        refresh_atual = tokens_atuais.get("refresh_token")

        if (
            access_atual
            and int(time.time()) < expires_at_atual
            and access_atual != tokens.get("access_token")
        ):
            print("[OAUTH DEBUG] Tokens já foram renovados por outra execução.")
            return {
                "access_token": access_atual,
                "refresh_token": refresh_atual,
                "expires_in": max(1, expires_at_atual - int(time.time()))
            }

        refresh_token = refresh_atual or refresh_token

        print("[OAUTH DEBUG] Enviando refresh_token ao Tiny:", resumo_token(refresh_token))
        print("[OAUTH DEBUG] client_id configurado:", bool(TINY_CLIENT_ID))
        print("[OAUTH DEBUG] client_secret configurado:", bool(TINY_CLIENT_SECRET))
        print("[OAUTH DEBUG] URL OAuth:", TINY_TOKEN_URL)

        response = requests.post(
            TINY_TOKEN_URL,
            data={
                "grant_type": "refresh_token",
                "client_id": TINY_CLIENT_ID,
                "client_secret": TINY_CLIENT_SECRET,
                "refresh_token": refresh_token
            },
            headers={
                "Accept": "application/json",
                "Content-Type": "application/x-www-form-urlencoded"
            },
            timeout=30
        )

        dados = resposta_json(response)

        print("[OAUTH DEBUG] Renovação OAuth Tiny HTTP:", response.status_code)
        print("[OAUTH DEBUG] Content-Type:", response.headers.get("Content-Type"))

        dados_log = dict(dados) if isinstance(dados, dict) else {"resposta": dados}
        if isinstance(dados_log, dict):
            if "access_token" in dados_log:
                dados_log["access_token"] = resumo_token(dados_log["access_token"])
            if "refresh_token" in dados_log:
                dados_log["refresh_token"] = resumo_token(dados_log["refresh_token"])
        print("[OAUTH DEBUG] Resposta resumida do Tiny:", dados_log)

        if not response.ok:
            texto_erro = json.dumps(dados, ensure_ascii=False).lower()
            invalid_grant = (
                "invalid_grant" in texto_erro
                or "token is not active" in texto_erro
                or "invalid_token" in texto_erro
            )

            if invalid_grant:
                print("[OAUTH DEBUG] O refresh_token foi rejeitado pelo Tiny.")
                print("[OAUTH DEBUG] Tokens NÃO serão apagados durante o diagnóstico.")
                raise TinyAPIError(
                    "O refresh token do Tiny não está mais ativo. É necessário autorizar novamente a aplicação no Tiny.",
                    401,
                    {
                        "resposta_tiny": dados,
                        "autorizacao": "/api/oauth/autorizar"
                    }
                )

            raise TinyAPIError(
                "Não foi possível renovar o access token.",
                response.status_code,
                dados
            )

        novo_access_token = dados.get("access_token")
        novo_refresh_token = dados.get("refresh_token") or refresh_token
        expires_in = dados.get("expires_in", 3600)

        print("[OAUTH DEBUG] Novo access_token:", resumo_token(novo_access_token))
        print("[OAUTH DEBUG] Novo refresh_token retornado:", resumo_token(dados.get("refresh_token")))
        print("[OAUTH DEBUG] Refresh token efetivamente salvo:", resumo_token(novo_refresh_token))
        print(
            "[OAUTH DEBUG] O refresh_token mudou:",
            fingerprint_token(novo_refresh_token) != fingerprint_token(refresh_token)
        )
        print("[OAUTH DEBUG] expires_in recebido:", expires_in)

        if not novo_access_token:
            raise TinyAPIError("Tiny não retornou novo access_token.", 502, dados)

        salvar_tokens(
            novo_access_token,
            novo_refresh_token,
            expires_in
        )

        print("[OAUTH DEBUG] Renovação concluída com sucesso.")
        return {
            "access_token": novo_access_token,
            "refresh_token": novo_refresh_token,
            "expires_in": expires_in
        }

    except requests.RequestException as e:
        print("[OAUTH DEBUG] Erro de comunicação durante o refresh:", repr(e))
        raise TinyAPIError(
            "Erro de comunicação durante a renovação do token.",
            None,
            str(e)
        )

    finally:
        if lock_adquirido:
            try:
                redis_delete(TINY_REFRESH_LOCK_KEY)
                print("[OAUTH DEBUG] Lock de renovação OAuth liberado.")
            except Exception as e:
                print("[OAUTH DEBUG] Não foi possível liberar o lock:", str(e))

def obter_access_token():

    tokens = carregar_tokens()


    if not tokens:

        raise TinyAPIError(

            "A aplicação ainda não foi autorizada no Tiny.",

            401,

            {
                "autorizacao":
                    "/api/oauth/autorizar"
            }
        )


    access_token = tokens.get(
        "access_token"
    )

    expires_at = int(
        tokens.get(
            "expires_at",
            0
        )
    )


    agora = int(
        time.time()
    )


    if (
        access_token
        and
        agora < expires_at
    ):

        return access_token


    print(
        "Access token expirado. Renovando..."
    )


    novos_tokens = renovar_access_token(
        tokens
    )


    return novos_tokens[
        "access_token"
    ]

def tiny_request(
    metodo,
    endpoint,
    **kwargs
):

    access_token = obter_access_token()


    response = requests.request(

        metodo,

        f"{TINY_API_URL}{endpoint}",

        headers=headers_tiny(
            access_token
        ),

        timeout=30,

        **kwargs
    )

    if response.status_code == 401:

        print(
            "Tiny retornou HTTP 401."
        )

        print(
            "Tentando renovar o access token..."
        )


        tokens = carregar_tokens()


        if not tokens:

            raise TinyAPIError(

                "Tokens OAuth não encontrados.",

                401,

                resposta_json(
                    response
                )
            )


        novos_tokens = renovar_access_token(
            tokens
        )


        response = requests.request(

            metodo,

            f"{TINY_API_URL}{endpoint}",

            headers=headers_tiny(

                novos_tokens[
                    "access_token"
                ]
            ),

            timeout=30,

            **kwargs
        )


    return response

def limpar_documento(
    valor
):

    if not valor:

        return ""


    return "".join(

        c

        for c in str(
            valor
        )

        if c.isdigit()
    )


def localizar_contato(
    cpf_cnpj,
    nome=None,
    busca_exaustiva=False
):
    documento = limpar_documento(cpf_cnpj)

    if not documento:
        raise TinyAPIError(
            "CPF/CNPJ do cliente não informado."
        )

    situacoes = [None, "B", "A", "I", "E"]
    vistos = set()

    for situacao in situacoes:
        params = {
            "cpfCnpj": documento,
            "limit": 100,
            "offset": 0
        }

        if situacao:
            params["situacao"] = situacao

        response = tiny_request(
            "GET",
            "/contatos",
            params=params
        )

        dados = resposta_json(response)

        print(
            "Consulta contato por CPF/CNPJ:",
            documento,
            "situação:",
            situacao or "todas",
            "HTTP:",
            response.status_code
        )

        if not response.ok:

            texto_erro = json.dumps(
                dados,
                ensure_ascii=False
            ).lower()

            contato_nao_encontrado = (
                "cpf/cnpj not found" in texto_erro
                or
                "cpf/cnpj não encontrado" in texto_erro
                or
                "cpf/cnpj nao encontrado" in texto_erro
                or
                "cpf ou cnpj não encontrado" in texto_erro
                or
                "cpf ou cnpj nao encontrado" in texto_erro
                or
                (
                    response.status_code == 404
                    and (
                        "not found" in texto_erro
                        or
                        "não encontrado" in texto_erro
                        or
                        "nao encontrado" in texto_erro
                    )
                )
            )

            if contato_nao_encontrado:
                print(
                    "CPF/CNPJ não encontrado no Tiny. "
                    "O contato será criado automaticamente."
                )

                return None

            if situacao:
                continue
            
            raise TinyAPIError(
                "Erro ao consultar contato no Tiny.",
                response.status_code,
                dados
            )

        contatos = dados.get(
            "itens",
            []
        )

        if not isinstance(contatos, list):
            contatos = []

        for contato in contatos:
            contato_id = contato.get("id")

            if contato_id in vistos:
                continue

            vistos.add(contato_id)

            documento_tiny = limpar_documento(
                contato.get("cpfCnpj")
            )

            if documento_tiny == documento:
                return contato
            
    if nome:
        response = tiny_request(
            "GET",
            "/contatos",
            params={
                "nome": nome,
                "limit": 100,
                "offset": 0
            }
        )

        dados = resposta_json(response)

        print(
            "Fallback consulta contato por nome:",
            nome,
            "HTTP:",
            response.status_code
        )

        if response.ok:
            contatos = dados.get(
                "itens",
                []
            )

            if isinstance(contatos, list):
                for contato in contatos:
                    documento_tiny = limpar_documento(
                        contato.get("cpfCnpj")
                    )

                    if documento_tiny == documento:
                        return contato

    if busca_exaustiva:
        limit = 100
        offset = 0
        total = None
        max_paginas = 1000

        for _ in range(max_paginas):

            response = tiny_request(
                "GET",
                "/contatos",
                params={
                    "limit": limit,
                    "offset": offset
                }
            )

            dados = resposta_json(response)

            print(
                "Busca exaustiva de contato:",
                "offset=",
                offset,
                "HTTP=",
                response.status_code
            )

            if not response.ok:
                raise TinyAPIError(
                    "Erro ao percorrer contatos do Tiny para localizar o CPF/CNPJ.",
                    response.status_code,
                    dados
                )

            contatos = dados.get(
                "itens",
                []
            )

            if not isinstance(contatos, list):
                contatos = []

            for contato in contatos:
                documento_tiny = limpar_documento(
                    contato.get("cpfCnpj")
                )

                if documento_tiny == documento:
                    print(
                        "Contato localizado na busca exaustiva. ID:",
                        contato.get("id")
                    )

                    return contato

            paginacao = dados.get(
                "paginacao",
                {}
            )

            if isinstance(paginacao, dict):
                try:
                    total = int(
                        paginacao.get("total")
                    )
                except (TypeError, ValueError):
                    total = None

            if not contatos:
                break

            offset += len(contatos)

            if total is not None and offset >= total:
                break

            if len(contatos) < limit and total is None:
                break

    return None

def criar_contato(dados_front):

    cliente = dados_front.get(
        "cliente",
        {}
    )

    endereco = dados_front.get(
        "endereco",
        {}
    )

    documento = limpar_documento(
        cliente.get("cpf_cnpj")
    )

    nome = (
        cliente.get("nome")
        or cliente.get("razao_social")
        or "Cliente da loja"
    )

    if not documento:
        raise TinyAPIError(
            "CPF/CNPJ do cliente não informado."
        )

    if len(documento) not in (11, 14):
        raise TinyAPIError(
            "CPF/CNPJ inválido. O documento deve conter 11 dígitos (CPF) "
            "ou 14 dígitos (CNPJ)."
        )

    if not nome:
        raise TinyAPIError(
            "Nome do cliente não informado."
        )

    codigo = f"{documento}"


    tipo_pessoa = "J" if len(documento) == 14 else "F"

    endereco_tiny = {
        "endereco": endereco.get("logradouro"),
        "numero": endereco.get("numero"),
        "complemento": endereco.get("complemento"),
        "bairro": endereco.get("bairro"),
        "municipio": endereco.get("cidade"),
        "cep": endereco.get("cep"),
        "uf": endereco.get("uf"),
        "pais": "Brasil"
    }

    endereco_tiny = {
        chave: valor
        for chave, valor in endereco_tiny.items()
        if valor not in [None, ""]
    }

    payload = {
        "nome": nome,
        "codigo": codigo,
        "tipoPessoa": tipo_pessoa,
        "cpfCnpj": documento,
        "email": dados_front.get("email"),
        "telefone": dados_front.get("telefone"),
        "endereco": endereco_tiny,
        "observacoesDoContato": ""
    }

    payload = {
        chave: valor
        for chave, valor in payload.items()
        if valor not in [None, ""]
    }

    print("")
    print("========================================")
    print("CRIANDO CONTATO NO TINY")
    print(json.dumps(payload, indent=2, ensure_ascii=False))
    print("========================================")

    response = tiny_request(
        "POST",
        "/contatos",
        json=payload
    )

    dados = resposta_json(response)

    print(
        "POST /contatos:",
        response.status_code
    )
    print(
        "Resposta criação contato:",
        dados
    )

    if not response.ok:

        texto_erro = json.dumps(
            dados,
            ensure_ascii=False
        ).lower()

        documento_duplicado = (
            response.status_code == 400
            and (
                "já existe" in texto_erro
                or "ja existe" in texto_erro
                or "already exists" in texto_erro
            )
            and (
                "cnpj" in texto_erro
                or "cpf" in texto_erro
            )
        )

        if documento_duplicado:

            print(
                "Tiny informou que o CPF/CNPJ já existe. "
                "Tentando localizar o contato existente..."
            )

            contato_existente = localizar_contato(
                documento,
                nome,
                busca_exaustiva=False
            )

            if contato_existente:

                contato_id = contato_existente.get(
                    "id"
                )

                if contato_id:
                    print(
                        "Contato existente recuperado após "
                        "erro de duplicidade. ID:",
                        contato_id
                    )

                    return {
                        "id": contato_id,
                        "nome": contato_existente.get(
                            "nome",
                            nome
                        ),
                        "cpfCnpj": documento,
                        "criado_agora": False,
                        "resposta": contato_existente,
                        "recuperado_apos_duplicidade": True
                    }

        raise TinyAPIError(
            "Tiny recusou a criação do contato.",
            response.status_code,
            dados
        )

    contato_id = None

    if isinstance(dados, dict):
        contato_id = dados.get("id")

        if not contato_id and isinstance(dados.get("data"), dict):
            contato_id = dados["data"].get("id")

    if not contato_id:
        raise TinyAPIError(
            "Tiny criou o contato, mas não retornou o ID.",
            502,
            dados
        )

    return {
        "id": contato_id,
        "nome": nome,
        "cpfCnpj": documento,
        "tipoPessoa": tipo_pessoa,
        "criado_agora": True,
        "resposta": dados
    }

def obter_ou_criar_contato(dados_front):

    cliente = dados_front.get(
        "cliente",
        {}
    )

    documento = limpar_documento(
        cliente.get("cpf_cnpj")
    )

    if not documento:
        raise TinyAPIError(
            "CPF/CNPJ do cliente não informado.",
            400
        )

    nome = (
        cliente.get("nome")
        or cliente.get("razao_social")
        or None
    )

    contato = localizar_contato(
        documento,
        nome,
        busca_exaustiva=False
    )

    if contato:

        contato_id = contato.get("id")

        if not contato_id:
            raise TinyAPIError(
                "Contato encontrado sem ID.",
                502,
                contato
            )

        print(
            "Contato encontrado no Tiny. ID:",
            contato_id
        )

        return {
            "id": contato_id,
            "criado_agora": False,
            "resposta": contato
        }

    print(
        "Contato não encontrado. Criando novo contato..."
    )

    return criar_contato(
        dados_front
    )

def obter_produto_por_id(produto_id):
    if not produto_id:
        return None

    response = tiny_request(
        "GET",
        f"/produtos/{produto_id}"
    )

    dados = resposta_json(response)

    print(
        "Consulta produto por ID",
        produto_id,
        "HTTP",
        response.status_code
    )

    if not response.ok:
        raise TinyAPIError(
            "Erro ao obter produto pelo ID.",
            response.status_code,
            dados
        )

    return dados

def ObterImagemDoProduto(produto_id):
    """Retorna a primeira URL de imagem/anexo válida do produto no Tiny."""

    if not produto_id:
        return None

    response = tiny_request(
        "GET",
        f"/produtos/{produto_id}/anexos"
    )

    dados = resposta_json(response)

    print(
        "Consulta imagens/anexos do produto",
        produto_id,
        "HTTP",
        response.status_code
    )

    if not response.ok:
        raise TinyAPIError(
            "Erro ao obter imagem do produto pelo ID.",
            response.status_code,
            dados
        )

    # A API retorna uma lista de anexos/imagens.
    if not isinstance(dados, list) or not dados:
        return None

    for anexo in dados:
        if not isinstance(anexo, dict):
            continue

        url = anexo.get("url")
        if url:
            return url

    return None


def carregar_imagem_produto(url, max_width=18 * mm, max_height=18 * mm):
    if not url:
        return None

    try:
        response = requests.get(
            url,
            timeout=10
        )
        response.raise_for_status()

        image_bytes = response.content
        if not image_bytes:
            return None

        reader = ImageReader(
            BytesIO(image_bytes)
        )
        largura_original, altura_original = reader.getSize()

        if not largura_original or not altura_original:
            return None

        escala = min(
            max_width / float(largura_original),
            max_height / float(altura_original)
        )

        largura = largura_original * escala
        altura = altura_original * escala

        imagem = Image(
            BytesIO(image_bytes),
            width=largura,
            height=altura
        )
        imagem.hAlign = "CENTER"

        return imagem

    except Exception as e:
        print(
            "Erro ao carregar imagem do produto:",
            e
        )
        return None


def localizar_produto_por_sku(
    sku
):

    if not sku:
        return None

    sku = str(sku).strip()

    if not sku:
        return None

    response = tiny_request(
        "GET",
        "/produtos",
        params={
            "codigo":
                sku,
            "limit":
                100,
            "offset":
                0
        }
    )

    dados = resposta_json(
        response
    )

    print(
        "Consulta produto SKU",
        sku,
        "HTTP",
        response.status_code
    )

    if not response.ok:
        raise TinyAPIError(
            "Erro ao consultar produto pelo SKU.",
            response.status_code,
            dados
        )

    produtos = dados.get(
        "itens",
        []
    )

    if not produtos:
        return None
    
    sku_normalizado = sku.casefold()

    for produto in produtos:
        sku_produto = produto.get(
            "sku"
        )

        if sku_produto is None:
            continue

        sku_produto_normalizado = str(
            sku_produto
        ).strip().casefold()

        if (
            sku_produto_normalizado
            and
            sku_produto_normalizado == sku_normalizado
        ):
            return produto

    return None

@app.route(
    "/api/testar-tiny",
    methods=["GET"]
)
def testar_tiny():

    try:

        response = tiny_request(

            "GET",

            "/contatos",

            params={

                "limit":
                    1,

                "offset":
                    0
            }
        )


        dados = resposta_json(
            response
        )


        if not response.ok:

            return jsonify({

                "erro":
                    "Token rejeitado pelo Tiny.",

                "status_tiny":
                    response.status_code,

                "resposta_tiny":
                    dados

            }), response.status_code


        return jsonify({

            "sucesso":
                True,

            "mensagem":
                "Autenticação com a API V3 funcionando.",

            "tiny":
                dados

        }), 200


    except TinyAPIError as e:

        return jsonify({

            "erro":
                e.mensagem,

            "status_tiny":
                e.status,

            "detalhes":
                e.resposta

        }), e.status or 500

@app.route(
    "/api/gerar-proposta",
    methods=["POST"]
)
def gerar_proposta():

    try:

        dados_front = request.get_json(
            silent=True
        )


        if not dados_front:

            return jsonify({

                "erro":
                    "JSON inválido ou vazio."

            }), 400

        cliente = dados_front.get(
            "cliente",
            {}
        )


        cpf_cnpj = cliente.get(
            "cpf_cnpj"
        )


        if not cpf_cnpj:

            return jsonify({

                "erro":
                    "CPF/CNPJ do cliente não informado."

            }), 400


        contato = obter_ou_criar_contato(
            dados_front
        )


        contato_id = contato.get(
            "id"
        )


        if not contato_id:

            raise TinyAPIError(
                "Não foi possível obter o ID do contato.",
                502,
                contato
            )

        carrinho = dados_front.get(
            "carrinho",
            []
        )


        if not carrinho:

            return jsonify({

                "erro":
                    "Carrinho vazio."

            }), 400


        itens_tiny = []

        for indice, item in enumerate(
            carrinho,
            start=1
        ):

            sku = item.get(
                "sku"
            )

            if sku is not None:
                sku = str(sku).strip()

            if not sku:

                return jsonify({

                    "erro":
                        "Produto sem SKU.",

                    "item":
                        indice,

                    "produto":
                        item

                }), 400


            produto = localizar_produto_por_sku(
                sku
            )


            if not produto:

                return jsonify({

                    "erro":
                        "Produto não encontrado no Tiny pelo SKU.",

                    "item":
                        indice,

                    "sku":
                        sku,

                    "nome_site":
                        (
                            item.get(
                                "nome"
                            )
                            or
                            item.get(
                                "descricao"
                            )
                        )

                }), 404


            produto_id = produto.get(
                "id"
            )


            if not produto_id:

                return jsonify({

                    "erro":
                        "Produto encontrado sem ID no Tiny.",

                    "produto":
                        produto

                }), 502


            quantidade = float(
                item.get(
                    "quantidade",
                    1
                )
            )


            preco = float(
                item.get(
                    "preco_unitario",
                    0
                )
            )


            item_tiny = {

                "produto": {

                    "id":
                        produto_id
                },

                "quantidade":
                    quantidade,

                "valorUnitario":
                    preco
            }

            produto_detalhado = obter_produto_por_id(
                produto_id
            )

            descricao_complementar = (
                produto_detalhado.get("descricaoComplementar")
                if isinstance(produto_detalhado, dict)
                else None
            )

            if descricao_complementar:
                item_tiny["descrComplementarOrc"] = html_para_texto(descricao_complementar)


            itens_tiny.append(
                item_tiny
            )

        introducao_proposta = (
            "Prezado cliente, seguem abaixo proposta comercial com "
            "pagamento à vista com desconto e nossos dados bancários:\n\n"
            "Segue nossos dados bancários:\n"
            "BRFER Comércio de Ferramentas LTDA\n"
            "CNPJ 40.954.410/0001-96\n"
            "Banco: 341 – Itaú\n"
            "Agência: 8811\n"
            "Conta Corrente: 99874-2\n\n"
            "Se preferir, o pagamento pode ser realizado via PIX, a chave "
            "é o nosso CNPJ"
        )

        resumo_carrinho = dados_front.get(
            "resumo_carrinho",
            {}
        )

        def valor_float(nome):
            try:
                return float(resumo_carrinho.get(nome, 0) or 0)
            except (TypeError, ValueError):
                return 0.0

        total_carrinho = valor_float("total")
        valor_avista = valor_float("avista")
        valor_parcela_3x = valor_float("parcela_3x")
        valor_parcela_12x = valor_float("parcela_12x")

        def dinheiro(valor):
            return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

        hoje = datetime.now(ZoneInfo("America/Sao_Paulo")).date()
        data_proximo_contato = hoje + timedelta(days=3)

        data_proposta = hoje.isoformat()
        data_proximo_contato_str = data_proximo_contato.isoformat()

        outros_itens_servicos = (
            dados_front.get(
                "outros_itens_servicos"
            )
            or
            (
                "Condições de pagamento\n"
                f"Total: {dinheiro(total_carrinho)}\n"
                f"Pagamento à vista com desconto: {dinheiro(valor_avista)}\n"
                f"3x de {dinheiro(valor_parcela_3x)} sem juros\n"
                f"12x de {dinheiro(valor_parcela_12x)} com juros no cartão."
            )
        )

        hoje = datetime.now(
            ZoneInfo("America/Sao_Paulo")
        ).date()

        data_proposta = hoje.isoformat()

        data_proximo_contato = (
            hoje + timedelta(days=3)
        ).isoformat()

        observacao_padrao = (
            dados_front.get(
                "observacoes"
            )
            or
            "Somos um E-COMMERCE, não reservamos estoque "
            "antes da aprovação do pagamento."
        )

        observacao_pagamento = (
            f"{observacao_padrao}\n\n"
            "Condições de pagamento:\n"
            f"Total: {dinheiro(total_carrinho)}\n"
            f"Pagamento à vista com desconto: {dinheiro(valor_avista)}\n"
            f"3x de {dinheiro(valor_parcela_3x)} sem juros\n"
            f"12x de {dinheiro(valor_parcela_12x)} com juros no cartão."
        )

        payload_tiny = {

            "contato": {

                "id":
                    contato_id
            },

            "itens":
                itens_tiny,

            "introducao":
                introducao_proposta,

            "data":
                data_proposta,

            "dataProximoContato":
                data_proximo_contato_str,

            "outrosItensServicos":
                outros_itens_servicos,

            "observacao":
                observacao_pagamento,
        }


        print("")
        print(
            "========================================"
        )

        print(
            "CRIANDO PROPOSTA NO TINY"
        )

        print(
            json.dumps(
                payload_tiny,
                indent=2,
                ensure_ascii=False
            )
        )

        print(
            "========================================"
        )

        response_post = tiny_request(

            "POST",

            "/orcamentos",

            json=payload_tiny
        )


        dados_criacao = resposta_json(
            response_post
        )


        print(
            "POST /orcamentos:",
            response_post.status_code
        )


        if not response_post.ok:

            return jsonify({

                "erro":
                    "Tiny recusou a criação da proposta.",

                "status_tiny":
                    response_post.status_code,

                "resposta_tiny":
                    dados_criacao

            }), response_post.status_code

        orcamento_id = None

        if isinstance(
            dados_criacao,
            dict
        ):

            orcamento_id = (

                dados_criacao.get(
                    "id"
                )

                or

                dados_criacao.get(
                    "idOrcamento"
                )
            )


        if not orcamento_id:

            return jsonify({

                "erro":
                    (
                        "Tiny respondeu sucesso, "
                        "mas não retornou o ID da proposta."
                    ),

                "resposta_tiny":
                    dados_criacao

            }), 502

        response_get = tiny_request(

            "GET",

            f"/orcamentos/{orcamento_id}"
        )


        dados_orcamento = resposta_json(
            response_get
        )


        print(
            "GET /orcamentos/",
            orcamento_id,
            ":",
            response_get.status_code
        )


        if response_get.ok:

            # O PDF é produzido localmente pelo backend, usando os dados
            # efetivamente persistidos no Tiny e os dados do formulário.
            pdf_buffer = gerar_pdf_proposta(
                dados_front,
                dados_orcamento,
                contato,
                orcamento_id
            )

            resposta = send_file(
                pdf_buffer,
                mimetype="application/pdf",
                as_attachment=False,
                download_name=f"proposta_comercial{orcamento_id}.pdf"
            )

            # Força abertura inline no navegador em vez de download automático.
            resposta.headers["Content-Disposition"] = (
                f'inline; filename="proposta_comercial{orcamento_id}.pdf"'
            )
            resposta.headers["X-Proposta-Id"] = str(orcamento_id)

            return resposta

        return jsonify({

            "sucesso":
                True,

            "id":
                orcamento_id,

            "contato": {
                "id": contato_id,
                "criado_agora": contato.get("criado_agora", False)
            },

            "criacao":
                dados_criacao,

            "erro_get":
                True,

            "status_get_tiny":
                response_get.status_code,

            "resposta_get_tiny":
                dados_orcamento

        }), 200


    except TinyAPIError as e:

        print("")
        print(
            "========================================"
        )

        print(
            "ERRO TINY"
        )

        print(
            "MENSAGEM:",
            e.mensagem
        )

        print(
            "STATUS:",
            e.status
        )

        print(
            "RESPOSTA:",
            e.resposta
        )

        print(
            "========================================"
        )


        return jsonify({

            "erro":
                e.mensagem,

            "status_tiny":
                e.status,

            "resposta_tiny":
                e.resposta

        }), e.status or 502


    except requests.RequestException as e:

        return jsonify({

            "erro":
                "Erro de comunicação com o Tiny.",

            "detalhes":
                str(e)

        }), 502


    except Exception as e:

        print(
            "ERRO INTERNO:",
            str(e)
        )


        return jsonify({

            "erro":
                "Erro interno no servidor.",

            "detalhes":
                str(e)

        }), 500

@app.route(
    "/api/obter-proposta/<int:id_proposta>",
    methods=["GET"]
)
def obter_proposta(
    id_proposta
):

    try:

        response = tiny_request(

            "GET",

            f"/orcamentos/{id_proposta}"
        )


        dados = resposta_json(
            response
        )


        if not response.ok:

            return jsonify({

                "erro":
                    "Falha ao obter o orçamento.",

                "status_tiny":
                    response.status_code,

                "resposta_tiny":
                    dados

            }), response.status_code


        return jsonify(
            dados
        ), 200


    except TinyAPIError as e:

        return jsonify({

            "erro":
                e.mensagem,

            "status_tiny":
                e.status,

            "resposta_tiny":
                e.resposta

        }), e.status or 502


    except Exception as e:

        return jsonify({

            "erro":
                "Erro interno.",

            "detalhes":
                str(e)

        }), 500


@app.route(
    "/api/oauth/testar-refresh",
    methods=["GET"]
)
def testar_refresh_oauth():
    inicio = int(time.time())
    print("#" * 80)
    print("[OAUTH TESTE] INÍCIO DO TESTE DE REFRESH")
    print("[OAUTH TESTE] Horário:", formatar_timestamp(inicio))
    print("#" * 80)

    try:
        tokens = carregar_tokens()
        if not tokens:
            return jsonify({
                "sucesso": False,
                "mensagem": "Nenhum token encontrado. Faça a autorização manual primeiro.",
                "autorizacao": "/api/oauth/autorizar"
            }), 401

        log_resumo_tokens("TOKENS ANTES DO TESTE", tokens)
        print("[OAUTH TESTE] Forçando refresh mesmo com access_token válido.")

        novos_tokens = renovar_access_token(tokens)
        tokens_depois = carregar_tokens()
        log_resumo_tokens("TOKENS DEPOIS DO TESTE", tokens_depois)

        fim = int(time.time())
        print("[OAUTH TESTE] SUCESSO. Duração:", fim - inicio, "segundos")

        return jsonify({
            "sucesso": True,
            "mensagem": "Refresh automático executado com sucesso. Consulte os logs da Vercel.",
            "access_token_recebido": bool(novos_tokens.get("access_token")),
            "refresh_token_recebido": bool(novos_tokens.get("refresh_token")),
            "expires_in": novos_tokens.get("expires_in"),
            "duracao_segundos": fim - inicio
        }), 200

    except TinyAPIError as e:
        print("[OAUTH TESTE] FALHA:", e.mensagem)
        return jsonify({
            "sucesso": False,
            "mensagem": e.mensagem,
            "detalhes": e.resposta,
            "autorizacao": "/api/oauth/autorizar"
        }), e.status or 500

    except Exception as e:
        print("[OAUTH TESTE] ERRO INESPERADO:", repr(e))
        return jsonify({
            "sucesso": False,
            "mensagem": "Erro inesperado durante o teste.",
            "detalhes": str(e)
        }), 500

@app.route("/api/oauth/renovar-agendado", methods=["GET"])
def renovar_oauth_agendado():

    inicio = int(time.time())

    print("")
    print("#" * 80)
    print("[OAUTH CRON] INÍCIO DA RENOVAÇÃO AGENDADA")
    print("[OAUTH CRON] Horário:", formatar_timestamp(inicio))
    print("#" * 80)

    cron_secret = os.getenv("CRON_SECRET")

    authorization = request.headers.get(
        "Authorization",
        ""
    )

    token_recebido = ""

    if authorization.startswith("Bearer "):
        token_recebido = authorization[7:]

    if not cron_secret:
        print(
            "[OAUTH CRON] ERRO: CRON_SECRET não configurado."
        )

        return jsonify({
            "sucesso": False,
            "mensagem": "CRON_SECRET não configurado no servidor."
        }), 500

    if token_recebido != cron_secret:
        print(
            "[OAUTH CRON] Acesso negado: segredo do cron inválido."
        )

        return jsonify({
            "sucesso": False,
            "mensagem": "Não autorizado."
        }), 401

    print(
        "[OAUTH CRON] Autenticação do cron validada."
    )

    try:
        tokens_antes = carregar_tokens()

        if not tokens_antes:
            print(
                "[OAUTH CRON] Nenhum token encontrado no Redis."
            )

            return jsonify({
                "sucesso": False,
                "mensagem": (
                    "Nenhum token OAuth armazenado. "
                    "É necessária uma autorização manual."
                ),
                "precisa_autorizar": True
            }), 401

        log_resumo_tokens(
            "TOKENS ANTES DA RENOVAÇÃO AGENDADA",
            tokens_antes
        )

        print(
            "[OAUTH CRON] Forçando renovação do access_token."
        )

        novos_tokens = renovar_access_token(
            tokens_antes
        )

        tokens_depois = carregar_tokens()

        if tokens_depois:
            log_resumo_tokens(
                "TOKENS DEPOIS DA RENOVAÇÃO AGENDADA",
                tokens_depois
            )

        fim = int(time.time())

        print("")
        print("#" * 80)
        print("[OAUTH CRON] RENOVAÇÃO CONCLUÍDA COM SUCESSO")
        print(
            "[OAUTH CRON] Duração:",
            fim - inicio,
            "segundos"
        )
        print("#" * 80)

        return jsonify({
            "sucesso": True,
            "mensagem": "Renovação OAuth agendada executada com sucesso.",
            "access_token_recebido": bool(
                novos_tokens.get("access_token")
            ),
            "refresh_token_recebido": bool(
                novos_tokens.get("refresh_token")
            ),
            "expires_in": novos_tokens.get("expires_in"),
            "duracao_segundos": fim - inicio
        }), 200

    except TinyAPIError as e:
        fim = int(time.time())

        print("")
        print("#" * 80)
        print("[OAUTH CRON] FALHA NA RENOVAÇÃO")
        print("[OAUTH CRON] Mensagem:", str(e))
        print(
            "[OAUTH CRON] Duração:",
            fim - inicio,
            "segundos"
        )
        print("#" * 80)

        return jsonify({
            "sucesso": False,
            "mensagem": str(e),
            "detalhes": getattr(e, "detalhes", None),
            "precisa_autorizar": True,
            "autorizacao": "/api/oauth/autorizar"
        }), getattr(e, "status_code", 500) or 500

    except Exception as e:
        print(
            "[OAUTH CRON] Erro inesperado:",
            repr(e)
        )

        return jsonify({
            "sucesso": False,
            "mensagem": "Erro inesperado durante a renovação agendada.",
            "detalhes": str(e)
        }), 500


@app.route(
    "/api/status",
    methods=["GET"]
)
def status():

    try:

        tokens = carregar_tokens()


        if not tokens:

            return jsonify({

                "autorizado":
                    False,

                "mensagem":
                    "Aplicação ainda não autorizada.",

                "autorizar":
                    "/api/oauth/autorizar"

            }), 200


        expires_at = int(
            tokens.get(
                "expires_at",
                0
            )
        )


        agora = int(
            time.time()
        )


        return jsonify({

            "autorizado":
                True,

            "access_token_valido":
                agora < expires_at,

            "tokens_armazenados":
                True,

            "mensagem":
                "Credenciais OAuth encontradas no Redis."

        }), 200


    except Exception as e:

        return jsonify({

            "erro":
                str(e)

        }), 500


@app.route(
    "/api/oauth/revogar",
    methods=["POST"]
)
def revogar_oauth():

    try:

        redis_delete(
            TINY_TOKEN_KEY
        )


        return jsonify({

            "sucesso":
                True,

            "mensagem":
                "Tokens removidos do Upstash Redis."

        }), 200


    except Exception as e:

        return jsonify({

            "erro":
                "Não foi possível remover os tokens.",

            "detalhes":
                str(e)

        }), 500

@app.route(
    "/",
    methods=["GET"]
)
def index():

    return jsonify({

        "status":
            "online",

        "servico":
            "Gerador de Propostas Comerciais",

        "api":
            "Olist ERP API V3",

        "endpoints": {

            "autorizar":
                "/api/oauth/autorizar",

            "callback":
                "/api/oauth/callback",

            "status":
                "/api/status",

            "testar_refresh":
                "/api/oauth/testar-refresh",

            "testar":
                "/api/testar-tiny",

            "gerar":
                "/api/gerar-proposta",

            "obter":
                "/api/obter-proposta/<id>",

            "revogar":
                "/api/oauth/revogar"
        }

    }), 200

@app.route(
    "/api/imprimir-proposta/<int:id_proposta>",
    methods=["POST"]
)
def imprimir_proposta(id_proposta):
    try:
        response = tiny_request(
            "POST",
            f"/orcamentos/{id_proposta}/imprimir"
        )
        
        dados = resposta_json(response)
        
        print(f"POST /orcamentos/{id_proposta}/imprimir:", response.status_code)

        if not response.ok:
            return jsonify({
                "erro": "Falha ao gerar o PDF da proposta.",
                "status_tiny": response.status_code,
                "resposta_tiny": dados
            }), response.status_code

        return jsonify(dados), 200

    except TinyAPIError as e:
        return jsonify({
            "erro": e.mensagem,
            "status_tiny": e.status,
            "resposta_tiny": e.resposta
        }), e.status or 502

    except Exception as e:
        return jsonify({
            "erro": "Erro interno ao tentar imprimir.",
            "detalhes": str(e)
        }), 500

@app.route("/api/listar_propostas_comerciais", methods=["GET"])
def listarPropostas(args = ''):
    try:
        response = tiny_request("GET", f"/orcamentos?{args}")

        dados = resposta_json(response)

        return dados

    except TinyAPIError as e:
        return jsonify({
            "erro": e.mensagem,
            "status_tiny": e.status,
            "resposta_tiny": e.resposta
        }), e.status or 502

    except Exception as e:
        return jsonify({
            "erro": "Erro interno ao tentar listar.",
            "detalhes": str(e)
        }), 500


def gerar_excel(titulos, linhasDoDF):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"

    borda = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000")
    )

    preenchimento_cabecalho = PatternFill(
        fill_type="solid",
        fgColor="000000"
    )

    fonte_cabecalho = Font(
        color="FFFFFF",
        bold=True
    )

    alinhamento = Alignment(
        horizontal="center",
        vertical="center"
    )

    formato_monetario = 'R$ #.##0,00'
    colunas_monetarias = {6, 12, 13}

    for coluna, titulo in enumerate(titulos, start=1):
        celula = ws.cell(
            row=1,
            column=coluna,
            value=titulo
        )

        celula.fill = preenchimento_cabecalho
        celula.font = fonte_cabecalho
        celula.border = borda
        celula.alignment = alinhamento

    for numero_linha, linha in enumerate(linhasDoDF, start=2):

        for numero_coluna, valor in enumerate(linha, start=1):

            celula = ws.cell(
                row=numero_linha,
                column=numero_coluna,
                value=valor
            )

            celula.border = borda
            celula.alignment = alinhamento

            indice = numero_coluna - 1 
            if indice in colunas_monetarias: 
                celula.number_format = 'R$ #.##0,00'

    for coluna in ws.columns:

        maior_tamanho = 0
        letra_coluna = coluna[0].column_letter

        for celula in coluna:

            if celula.value is not None:
                tamanho = len(str(celula.value))

                if celula.number_format == formato_monetario: 
                    tamanho += 3

                if tamanho > maior_tamanho:
                    maior_tamanho = tamanho

        largura = maior_tamanho + 2

        if largura < 10:
            largura = 10

        ws.column_dimensions[letra_coluna].width = largura

    ws.row_dimensions[1].height = 24
    arquivo = BytesIO()

    wb.save(arquivo)

    arquivo.seek(0)

    return arquivo

@app.route("/api/listarParaOSite", methods=["GET"])
def listarParaOSite():

    def getVendedor(saudacao, responsavel):
        if 'rose' in saudacao or 'rose' in responsavel:
            return 'Rose'
        elif 'amandha' in saudacao or 'amandha' in responsavel:
            return 'Amandha'
        elif 'jadhy' in saudacao or 'jadhy' in responsavel:
            return 'Jadhy'
        else:
            return ''

    data_inicio = request.args.get("data_inicio")
    data_fim = request.args.get("data_fim")
    vendedorGetParamether = request.args.get("vendedor")
    situacoes = request.args.getlist("situacoes")

    #TODO adicionar tratamento para o situacoes
    arrayPropostas = listarPropostas(f"dataInicio={data_inicio}&data_fim={data_fim}")

    titulos = ['Nº Da Proposta', 'Data', 'Data Prox Contato', 'Vendedor', 'Situação', 'Produto', 
               'Valor', 'Nome Cliente', 'Aos Cuidados', 'Fone', 'Celular', 'E-mail', 'Desconto', 'Frete']

    linhasDoDF = []
    itens = arrayPropostas.get("itens", [])

    for i in range(len(itens)):
        IdProposta = itens[i].get("id")
        requestPropostaMomentanea = tiny_request("GET", f"/orcamentos/{IdProposta}")
        requestPropostaMomentanea = resposta_json(requestPropostaMomentanea)

        aux1 = requestPropostaMomentanea.get("assinatura").get("saudacao", '').lower()
        aux2 = requestPropostaMomentanea.get("assinatura").get("responsavel", '').lower()

        vendedor = getVendedor(aux1, aux2)

        idContato = requestPropostaMomentanea.get("contato").get("id")
        contato = tiny_request("GET", f"/contatos/{idContato}")
        contato = resposta_json(contato)

        ProdutosProposta = requestPropostaMomentanea.get("itens", [])

        for j in range(len(ProdutosProposta)):
            line = ['N/A']*14
            line[0]  = itens[i].get("numeroProposta")
            line[1]  = itens[i].get("data")
            line[2]  = itens[i].get("dataProximoContato", "")
            line[3]  = vendedor
            line[4]  = itens[i].get("situacao", [])
            line[5]  = ProdutosProposta[j].get("produto").get("descricao", "")
            line[6]  = float(ProdutosProposta[j].get("valorUnitario")) * int(ProdutosProposta[j].get("quantidade"))
            try:
                line[7]  = contato.get("nome", "")
                line[8]  = contato.get("observacoesDoContato", "")
                line[9]  = contato.get("telefone", "")
                line[10] = contato.get("celular", "")
                line[11] = contato.get("email", "")
            except:
                pass
            line[12] = requestPropostaMomentanea.get("extras").get("desconto", 0)
            line[13] = requestPropostaMomentanea.get("extras").get("frete", 0)

            linhasDoDF.append(line)

    arquivo = gerar_excel(titulos, linhasDoDF)

    return send_file(
        arquivo,
        as_attachment=True,
        download_name="FiltroPropostas.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    #return jsonify({"titulos": titulos, "corpo": linhasDoDF})